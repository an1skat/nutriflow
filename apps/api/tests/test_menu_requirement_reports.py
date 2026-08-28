from datetime import date as Date
from decimal import Decimal
from io import BytesIO

import openpyxl
from beanie import PydanticObjectId
from fastapi.testclient import TestClient

from app.core.config import get_settings
from app.modules.menu_requirements.reporting import _build_calendar_month


def login(client: TestClient, identifier: str, password: str) -> None:
    response = client.post(
        "/api/v1/auth/login",
        json={"identifier": identifier, "password": password},
    )
    assert response.status_code == 204


def csrf_headers(client: TestClient) -> dict[str, str]:
    token = client.cookies.get(get_settings().csrf_cookie_name)
    assert token is not None
    return {"X-CSRF-Token": token}


def create_confirmed_dish(
    client: TestClient,
    *,
    ingredient_name: str = "Морква",
    card_number: str = "REQ-R-1",
    net_amount: str = "20.25",
    gross_amount: str = "25",
) -> tuple[str, str, str]:
    ingredient_response = client.post(
        "/api/v1/recipes/ingredients",
        json={"name": ingredient_name, "unit": "g"},
        headers=csrf_headers(client),
    )
    assert ingredient_response.status_code == 201
    ingredient_id = ingredient_response.json()["id"]

    card_response = client.post(
        "/api/v1/recipes/dish-cards",
        json={"card_number": card_number, "name": "Овочевий суп"},
        headers=csrf_headers(client),
    )
    assert card_response.status_code == 201
    card_id = card_response.json()["id"]
    variant_id = PydanticObjectId()

    version_response = client.post(
        f"/api/v1/recipes/dish-cards/{card_id}/versions",
        json={
            "portion_variants": [
                {
                    "id": str(variant_id),
                    "age_group": "6-11",
                    "output_grams": "200",
                    "nutrition": {},
                }
            ],
            "ingredient_amounts": [
                {
                    "ingredient_id": ingredient_id,
                    "ingredient_name_snapshot": ingredient_name,
                    "gross_amount": gross_amount,
                    "net_amount": net_amount,
                    "unit": "g",
                    "portion_variant_id": str(variant_id),
                }
            ],
        },
        headers=csrf_headers(client),
    )
    assert version_response.status_code == 201
    version_id = version_response.json()["id"]

    confirm_response = client.post(
        f"/api/v1/recipes/dish-card-versions/{version_id}/confirm",
        headers=csrf_headers(client),
    )
    assert confirm_response.status_code == 200
    return ingredient_id, card_id, str(variant_id)


def publish_school_menu(
    client: TestClient,
    *,
    school_id: str,
    card_id: str,
    variant_id: str,
    days: list[tuple[str, str]],
    starts_on: str = "2026-07-06",
    ends_on: str = "2026-07-10",
) -> dict:
    create_response = client.post(
        "/api/v1/menus/weekly",
        json={
            "title": "Меню для агрегатів",
            "meal_type": "lunch",
            "starts_on": starts_on,
            "ends_on": ends_on,
            "days": [
                {
                    "weekday": weekday,
                    "date": service_date,
                    "items": [
                        {
                            "position": 1,
                            "kind": "dish_card",
                            "recipe_card_number": "REQ-R-1",
                            "dish_card_id": card_id,
                            "name": "Овочевий суп",
                            "allergen_codes": [],
                            "portions": [
                                {
                                    "age_group": "6-11",
                                    "yield_amount": "200",
                                    "dish_card_portion_variant_id": variant_id,
                                    "nutrition": {},
                                }
                            ],
                        }
                    ],
                }
                for weekday, service_date in days
            ],
        },
        headers=csrf_headers(client),
    )
    assert create_response.status_code == 201
    source_id = create_response.json()["id"]

    publish_response = client.post(
        f"/api/v1/menus/weekly/{source_id}/publish",
        json={"school_ids": [school_id]},
        headers=csrf_headers(client),
    )
    assert publish_response.status_code == 200
    school_menu_id = publish_response.json()["created_menu_ids"][0]

    menu_response = client.get(f"/api/v1/menus/weekly/{school_menu_id}")
    assert menu_response.status_code == 200
    return menu_response.json()


def set_school_menu_counts(
    client: TestClient,
    *,
    menu: dict,
    group_id: str,
    counts_by_weekday: dict[str, int],
) -> dict:
    days = menu["days"]
    for day in days:
        day["items"][0]["servings"] = [
            {
                "school_group_id": group_id,
                "age_group": "6-11",
                "children_count": counts_by_weekday[day["weekday"]],
            }
        ]

    response = client.patch(
        f"/api/v1/menus/weekly/{menu['id']}",
        json={"days": days, "revision": menu["revision"]},
        headers=csrf_headers(client),
    )
    assert response.status_code == 200
    return response.json()


def generate_requirement(
    client: TestClient,
    *,
    menu_id: str,
    weekday: str,
    service_date: str,
) -> dict:
    response = client.post(
        "/api/v1/menu-requirements/generate",
        json={
            "weekly_menu_id": menu_id,
            "weekday": weekday,
            "service_date": service_date,
        },
        headers=csrf_headers(client),
    )
    assert response.status_code == 200
    return response.json()["items"][0]


def assign_schools_to_community(
    client: TestClient,
    *school_ids: PydanticObjectId,
) -> None:
    for school_id in school_ids:
        response = client.patch(
            f"/api/v1/admin/schools/{school_id}",
            json={"community": "obukhivska"},
            headers=csrf_headers(client),
        )
        assert response.status_code == 200


def test_week_and_custom_range_reports_aggregate_daily_requirements(
    seeded_client,
) -> None:
    client, identities = seeded_client
    login(client, identities.admin.username, identities.admin_password)
    ingredient_id, card_id, variant_id = create_confirmed_dish(client)
    menu = publish_school_menu(
        client,
        school_id=str(identities.own_school.id),
        card_id=card_id,
        variant_id=variant_id,
        days=[
            ("monday", "2026-07-06"),
            ("tuesday", "2026-07-07"),
            ("wednesday", "2026-07-08"),
        ],
        starts_on="2026-07-06",
        ends_on="2026-07-08",
    )

    login(client, identities.school_user.username, identities.school_user_password)
    group_id = str(identities.own_school.groups[0].id)
    menu = set_school_menu_counts(
        client,
        menu=menu,
        group_id=group_id,
        counts_by_weekday={"monday": 3, "tuesday": 4, "wednesday": 2},
    )
    generate_requirement(
        client,
        menu_id=menu["id"],
        weekday="monday",
        service_date="2026-07-06",
    )
    generate_requirement(
        client,
        menu_id=menu["id"],
        weekday="wednesday",
        service_date="2026-07-08",
    )

    login(client, identities.admin.username, identities.admin_password)
    report_params = {
        "school_id": str(identities.own_school.id),
        "date_from": "2026-07-06",
        "date_to": "2026-07-08",
        "granularity": "week",
        "meal_type": "lunch",
        "school_group_id": group_id,
    }
    range_params = {**report_params, "granularity": "range"}
    incomplete_range_response = client.get(
        "/api/v1/menu-requirements/report",
        params=range_params,
    )
    assert incomplete_range_response.status_code == 400
    assert incomplete_range_response.json()["detail"] == (
        "Menu requirement report cannot be generated; missing dates: 2026-07-07"
    )

    response = client.get("/api/v1/menu-requirements/report", params=report_params)

    assert response.status_code == 200
    report = response.json()
    assert report["status"] == "missing"
    assert report["missing_dates"] == ["2026-07-07"]
    assert report["stale_dates"] == []

    group = report["groups"][0]
    assert group["school_group_id"] == group_id
    assert group["dishes"][0]["children_count_total"] == 5

    carrot = next(row for row in group["ingredient_rows"] if row["ingredient_id"] == ingredient_id)
    cell = carrot["cells"][0]
    assert Decimal(cell["net_per_person_g"]) == Decimal("40.50")
    assert Decimal(cell["gross_per_person_g"]) == Decimal("50")
    assert Decimal(cell["issue_total_raw_g"]) == Decimal("101.25")
    assert cell["issue_total_rounded_g"] == 102
    assert Decimal(cell["gross_issue_total_raw_g"]) == Decimal("125")
    assert cell["gross_issue_total_rounded_g"] == 125
    assert Decimal(carrot["issue_total_raw_g"]) == Decimal("101.25")
    assert carrot["issue_total_rounded_g"] == 102
    assert Decimal(carrot["gross_per_person_total_g"]) == Decimal("50")
    assert Decimal(carrot["gross_issue_total_raw_g"]) == Decimal("125")
    assert carrot["gross_issue_total_rounded_g"] == 125

    breakdown = cell["breakdown"]
    assert [item["service_date"] for item in breakdown] == [
        "2026-07-06",
        "2026-07-07",
        "2026-07-08",
    ]
    assert breakdown[0]["status"] == "complete"
    assert breakdown[0]["children_count"] == 3
    assert Decimal(breakdown[0]["issue_total_raw_g"]) == Decimal("60.75")
    assert Decimal(breakdown[0]["gross_per_person_g"]) == Decimal("25")
    assert Decimal(breakdown[0]["gross_issue_total_raw_g"]) == Decimal("75")
    assert breakdown[0]["gross_issue_total_rounded_g"] == 75
    assert breakdown[1]["status"] == "missing"
    assert breakdown[1]["requirement_id"] is None
    assert breakdown[1]["net_per_person_g"] is None
    assert breakdown[1]["gross_per_person_g"] is None
    assert breakdown[2]["status"] == "complete"
    assert breakdown[2]["children_count"] == 2

    owner_export_response = client.get(
        "/api/v1/menu-requirements/report/export.xlsx",
        params=report_params,
    )
    assert owner_export_response.status_code == 200
    workbook = openpyxl.load_workbook(BytesIO(owner_export_response.content))
    sheet = workbook.active
    assert sheet["A1"].value == "МЕНЮ-ВИМОГА"
    assert sheet["B2"].value == identities.own_school.name
    assert any(cell.value == "Морква" for row in sheet.iter_rows() for cell in row)

    login(client, identities.school_user.username, identities.school_user_password)
    school_export_response = client.get(
        "/api/v1/menu-requirements/report/export.xlsx",
        params=report_params,
    )
    assert school_export_response.status_code == 400
    assert school_export_response.json()["detail"] == (
        "Weekly menu requirement report requires complete menu requirements for all five weekdays"
    )

    generate_requirement(
        client,
        menu_id=menu["id"],
        weekday="tuesday",
        service_date="2026-07-07",
    )
    complete_range_response = client.get(
        "/api/v1/menu-requirements/report",
        params=range_params,
    )
    assert complete_range_response.status_code == 200
    complete_range = complete_range_response.json()
    assert complete_range["granularity"] == "range"
    assert complete_range["status"] == "complete"
    assert complete_range["missing_dates"] == []
    assert complete_range["stale_dates"] == []
    assert complete_range["groups"][0]["dishes"][0]["children_count_total"] == 9


def test_report_marks_generated_day_stale_after_daily_menu_changes(
    seeded_client,
) -> None:
    client, identities = seeded_client
    login(client, identities.admin.username, identities.admin_password)
    _ingredient_id, card_id, variant_id = create_confirmed_dish(client)
    menu = publish_school_menu(
        client,
        school_id=str(identities.own_school.id),
        card_id=card_id,
        variant_id=variant_id,
        days=[("monday", "2026-07-06")],
    )

    login(client, identities.school_user.username, identities.school_user_password)
    group_id = str(identities.own_school.groups[0].id)
    menu = set_school_menu_counts(
        client,
        menu=menu,
        group_id=group_id,
        counts_by_weekday={"monday": 3},
    )
    requirement = generate_requirement(
        client,
        menu_id=menu["id"],
        weekday="monday",
        service_date="2026-07-06",
    )

    menu["days"][0]["items"][0]["servings"][0]["children_count"] = 7
    stale_response = client.patch(
        f"/api/v1/menus/weekly/{menu['id']}",
        json={"days": menu["days"], "revision": menu["revision"]},
        headers=csrf_headers(client),
    )
    assert stale_response.status_code == 200

    login(client, identities.admin.username, identities.admin_password)
    response = client.get(
        "/api/v1/menu-requirements/report",
        params={
            "school_id": str(identities.own_school.id),
            "date_from": "2026-07-06",
            "date_to": "2026-07-06",
            "granularity": "day",
            "meal_type": "lunch",
            "school_group_id": group_id,
        },
    )

    assert response.status_code == 200
    report = response.json()
    assert report["status"] == "stale"
    assert report["stale_dates"] == ["2026-07-06"]
    cell = report["groups"][0]["ingredient_rows"][0]["cells"][0]
    assert cell["breakdown"][0]["status"] == "stale"
    assert cell["breakdown"][0]["requirement_id"] == requirement["id"]


def test_calendar_uses_full_workweek_across_month_boundary() -> None:
    workweek = {
        Date(2026, 9, 28),
        Date(2026, 9, 29),
        Date(2026, 9, 30),
        Date(2026, 10, 1),
        Date(2026, 10, 2),
    }

    september = _build_calendar_month(
        2026,
        9,
        expected_dates=workweek,
        generated_dates=workweek,
        stale_dates=set(),
    )
    october = _build_calendar_month(
        2026,
        10,
        expected_dates=workweek,
        generated_dates=workweek,
        stale_dates=set(),
    )

    assert september.generated_days == 3
    assert october.generated_days == 2

    september_week = september.weeks[-1]
    october_week = october.weeks[0]
    assert (september_week.date_from, september_week.date_to) == (
        Date(2026, 9, 28),
        Date(2026, 10, 2),
    )
    assert (october_week.date_from, october_week.date_to) == (
        Date(2026, 9, 28),
        Date(2026, 10, 2),
    )
    assert september_week.generated_days == 5
    assert october_week.generated_days == 5

    august = _build_calendar_month(
        2026,
        8,
        expected_dates=set(),
        generated_dates=set(),
        stale_dates=set(),
    )
    assert (august.weeks[0].date_from, august.weeks[0].date_to) == (
        Date(2026, 8, 3),
        Date(2026, 8, 7),
    )


def test_calendar_treats_period_without_requirements_as_empty() -> None:
    expected_dates = {Date(2026, 8, day) for day in range(17, 22)}

    august = _build_calendar_month(
        2026,
        8,
        expected_dates=expected_dates,
        generated_dates=set(),
        stale_dates=set(),
    )

    week = next(item for item in august.weeks if item.date_from == Date(2026, 8, 17))
    assert august.missing_days == 0
    assert week.missing_days == 0


def test_calendar_keeps_an_existing_weekend_requirement_visible() -> None:
    generated_date = Date(2026, 7, 12)

    july = _build_calendar_month(
        2026,
        7,
        expected_dates={generated_date},
        generated_dates={generated_date},
        stale_dates=set(),
    )

    week = next(
        item for item in july.weeks if generated_date in {day.service_date for day in item.days}
    )
    assert july.generated_days == 1
    assert week.date_from == Date(2026, 7, 6)
    assert week.date_to == generated_date
    assert week.generated_days == 1


def test_calendar_access_and_workweek_block(seeded_client) -> None:
    client, identities = seeded_client
    login(client, identities.admin.username, identities.admin_password)
    _ingredient_id, card_id, variant_id = create_confirmed_dish(client)
    menu = publish_school_menu(
        client,
        school_id=str(identities.own_school.id),
        card_id=card_id,
        variant_id=variant_id,
        days=[("friday", "2026-07-31")],
        starts_on="2026-07-27",
        ends_on="2026-07-31",
    )

    login(client, identities.school_user.username, identities.school_user_password)
    group_id = str(identities.own_school.groups[0].id)
    menu = set_school_menu_counts(
        client,
        menu=menu,
        group_id=group_id,
        counts_by_weekday={"friday": 6},
    )
    generate_requirement(
        client,
        menu_id=menu["id"],
        weekday="friday",
        service_date="2026-07-31",
    )

    login(client, identities.lower_admin.username, identities.lower_admin_password)
    denied_response = client.get(
        "/api/v1/menu-requirements/calendar",
        params={"school_id": str(identities.other_school.id), "year": 2026},
    )
    assert denied_response.status_code == 403

    own_response = client.get(
        "/api/v1/menu-requirements/calendar",
        params={
            "school_id": str(identities.own_school.id),
            "year": 2026,
            "meal_type": "lunch",
            "school_group_id": group_id,
        },
    )
    assert own_response.status_code == 200
    july = next(month for month in own_response.json()["months"] if month["month"] == 7)
    assert july["status"] == "complete"
    assert july["generated_days"] == 1
    assert july["weeks"][-1]["date_from"] == "2026-07-27"
    assert july["weeks"][-1]["date_to"] == "2026-07-31"
    assert july["weeks"][-1]["generated_days"] == 1

    login(client, identities.admin.username, identities.admin_password)
    create_technologist_response = client.post(
        "/api/v1/admin/admins",
        json={
            "username": "calendar.tech",
            "email": "calendar.tech@example.com",
            "password": "tech-password-123",
            "role": "TECHNOLOGIST",
            "permissions": [],
        },
        headers=csrf_headers(client),
    )
    assert create_technologist_response.status_code == 201

    login(client, "calendar.tech", "tech-password-123")
    technologist_response = client.get(
        "/api/v1/menu-requirements/calendar",
        params={"school_id": str(identities.other_school.id), "year": 2026},
    )
    assert technologist_response.status_code == 200


def test_community_scope_access_is_all_or_nothing(seeded_client) -> None:
    client, identities = seeded_client

    login(client, identities.admin.username, identities.admin_password)
    assign_schools_to_community(
        client,
        identities.own_school.id,
        identities.other_school.id,
    )

    owner_response = client.get("/api/v1/menu-requirements/communities")
    assert owner_response.status_code == 200
    assert owner_response.json() == [
        {
            "community": "obukhivska",
            "community_name": "Обухівська громада",
            "school_count": 2,
        }
    ]

    invalid_community_response = client.get(
        "/api/v1/menu-requirements/communities/unknown/calendar",
        params={"year": 2026},
    )
    assert invalid_community_response.status_code == 422

    login(
        client,
        identities.lower_admin.username,
        identities.lower_admin_password,
    )
    lower_admin_list_response = client.get("/api/v1/menu-requirements/communities")
    assert lower_admin_list_response.status_code == 200
    assert lower_admin_list_response.json() == []

    lower_admin_calendar_response = client.get(
        "/api/v1/menu-requirements/communities/obukhivska/calendar",
        params={"year": 2026},
    )
    assert lower_admin_calendar_response.status_code == 403

    login(
        client,
        identities.school_user.username,
        identities.school_user_password,
    )
    school_user_response = client.get("/api/v1/menu-requirements/communities")
    assert school_user_response.status_code == 403


def test_community_calendar_report_and_export_aggregate_schools(
    seeded_client,
) -> None:
    client, identities = seeded_client
    other_school_password = "other-school-password-123"

    login(client, identities.admin.username, identities.admin_password)
    assign_schools_to_community(
        client,
        identities.own_school.id,
        identities.other_school.id,
    )

    other_user_response = client.post(
        f"/api/v1/admin/schools/{identities.other_school.id}/users",
        json={
            "username": "other.school.user",
            "email": "other.school.user@example.com",
            "password": other_school_password,
        },
        headers=csrf_headers(client),
    )
    assert other_user_response.status_code == 201

    ingredient_id, card_id, variant_id = create_confirmed_dish(client)
    own_menu = publish_school_menu(
        client,
        school_id=str(identities.own_school.id),
        card_id=card_id,
        variant_id=variant_id,
        days=[("monday", "2026-07-06")],
        starts_on="2026-07-06",
        ends_on="2026-07-06",
    )
    other_menu = publish_school_menu(
        client,
        school_id=str(identities.other_school.id),
        card_id=card_id,
        variant_id=variant_id,
        days=[("monday", "2026-07-06")],
        starts_on="2026-07-06",
        ends_on="2026-07-06",
    )

    login(
        client,
        identities.school_user.username,
        identities.school_user_password,
    )
    own_menu = set_school_menu_counts(
        client,
        menu=own_menu,
        group_id=str(identities.own_school.groups[0].id),
        counts_by_weekday={"monday": 3},
    )
    generate_requirement(
        client,
        menu_id=own_menu["id"],
        weekday="monday",
        service_date="2026-07-06",
    )

    login(client, "other.school.user", other_school_password)
    other_menu = set_school_menu_counts(
        client,
        menu=other_menu,
        group_id=str(identities.other_school.groups[0].id),
        counts_by_weekday={"monday": 4},
    )

    login(client, identities.admin.username, identities.admin_password)
    calendar_response = client.get(
        "/api/v1/menu-requirements/communities/obukhivska/calendar",
        params={"year": 2026, "meal_type": "lunch"},
    )
    assert calendar_response.status_code == 200

    calendar = calendar_response.json()
    assert calendar["community"] == "obukhivska"
    assert calendar["community_name"] == "Обухівська громада"
    assert calendar["school_count"] == 2

    july = next(month for month in calendar["months"] if month["month"] == 7)
    monday = next(
        day
        for week in july["weeks"]
        for day in week["days"]
        if day["service_date"] == "2026-07-06"
    )
    assert monday["expected_requirements"] == 2
    assert monday["generated_requirements"] == 1
    assert monday["missing_requirements"] == 1
    assert monday["status"] == "missing"
    assert july["generated_days"] == 1
    assert july["missing_days"] == 1

    day_report_params = {
        "date_from": "2026-07-06",
        "date_to": "2026-07-06",
        "granularity": "day",
        "meal_type": "lunch",
    }
    report_response = client.get(
        "/api/v1/menu-requirements/communities/obukhivska/report",
        params=day_report_params,
    )
    assert report_response.status_code == 200

    report = report_response.json()
    assert report["status"] == "missing"
    assert report["missing_dates"] == ["2026-07-06"]
    assert report["school_count"] == 2

    group = report["groups"][0]
    assert group["group_key"] == "age-group:6-11"
    assert group["school_group_id"] is None
    assert group["school_group_name"] == "6-11"
    assert group["dishes"][0]["children_count_total"] == 3

    ingredient = next(
        row
        for row in group["ingredient_rows"]
        if row["ingredient_id"] == ingredient_id
    )
    cell = ingredient["cells"][0]
    breakdown_by_school = {
        item["school_name"]: item
        for item in cell["breakdown"]
    }
    assert set(breakdown_by_school) == {
        identities.own_school.name,
        identities.other_school.name,
    }
    assert breakdown_by_school[identities.own_school.name]["status"] == "complete"
    assert breakdown_by_school[identities.own_school.name]["children_count"] == 3
    assert breakdown_by_school[identities.other_school.name]["status"] == "missing"
    assert breakdown_by_school[identities.other_school.name]["requirement_id"] is None

    login(client, "other.school.user", other_school_password)
    generate_requirement(
        client,
        menu_id=other_menu["id"],
        weekday="monday",
        service_date="2026-07-06",
    )

    login(client, identities.admin.username, identities.admin_password)
    range_report_params = {**day_report_params, "granularity": "range"}
    complete_response = client.get(
        "/api/v1/menu-requirements/communities/obukhivska/report",
        params=range_report_params,
    )
    assert complete_response.status_code == 200

    complete_report = complete_response.json()
    assert complete_report["status"] == "complete"
    assert complete_report["missing_dates"] == []
    assert complete_report["stale_dates"] == []

    complete_group = complete_report["groups"][0]
    assert complete_group["dishes"][0]["children_count_total"] == 7
    complete_ingredient = next(
        row
        for row in complete_group["ingredient_rows"]
        if row["ingredient_id"] == ingredient_id
    )
    complete_cell = complete_ingredient["cells"][0]
    assert Decimal(complete_cell["issue_total_raw_g"]) == Decimal("141.75")
    assert {item["status"] for item in complete_cell["breakdown"]} == {"complete"}
    assert {item["school_name"] for item in complete_cell["breakdown"]} == {
        identities.own_school.name,
        identities.other_school.name,
    }

    for granularity, date_to in (
        ("week", "2026-07-12"),
        ("month", "2026-07-31"),
    ):
        period_response = client.get(
            "/api/v1/menu-requirements/communities/obukhivska/report",
            params={
                "date_from": "2026-07-06",
                "date_to": date_to,
                "granularity": granularity,
                "meal_type": "lunch",
            },
        )
        assert period_response.status_code == 200
        assert period_response.json()["granularity"] == granularity

    export_response = client.get(
        "/api/v1/menu-requirements/communities/obukhivska/report/export.xlsx",
        params=range_report_params,
    )
    assert export_response.status_code == 200
    assert export_response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    workbook = openpyxl.load_workbook(BytesIO(export_response.content))
    sheet = workbook.active
    assert sheet["A2"].value == "Громада"
    assert sheet["B2"].value == "Обухівська громада"
    assert sheet["A3"].value == "Кількість шкіл"
    assert sheet["B3"].value == 2
    assert any(cell.value == "Морква" for row in sheet.iter_rows() for cell in row)
