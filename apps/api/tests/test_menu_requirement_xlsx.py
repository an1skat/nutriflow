from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from beanie import PydanticObjectId

from app.modules.identity.models import AgeGroup
from app.modules.menu_requirements.models import (
    MenuRequirement,
    MenuRequirementCell,
    MenuRequirementDish,
    MenuRequirementIngredientRow,
)
from app.modules.menu_requirements.schemas import MenuRequirementReportResponse
from app.modules.menu_requirements.xlsx import (
    build_menu_requirement_report_workbook,
    build_menu_requirement_workbook,
)
from app.modules.menus.models import MealType, Weekday

pytestmark = pytest.mark.no_clean_database


def test_builds_daily_menu_requirement_workbook() -> None:
    dish_id = PydanticObjectId()
    now = datetime.now(UTC)
    requirement = MenuRequirement.model_construct(
        school_id=PydanticObjectId(),
        weekly_menu_id=PydanticObjectId(),
        menu_title="Меню на тиждень",
        meal_type=MealType.LUNCH,
        weekday=Weekday.MONDAY,
        service_date=date(2026, 7, 6),
        school_group_id=PydanticObjectId(),
        school_group_name="1-А",
        age_group=AgeGroup.SIX_TO_ELEVEN,
        dishes=[
            MenuRequirementDish(
                menu_item_id=dish_id,
                position=1,
                kind="dish_card",
                name="Овочевий суп",
                yield_amount="200",
                children_count=3,
            )
        ],
        ingredient_rows=[
            MenuRequirementIngredientRow(
                key="ingredient:carrot",
                ingredient_name="Морква",
                cells=[
                    MenuRequirementCell(
                        menu_item_id=dish_id,
                        net_per_person_g=Decimal("20.25"),
                    )
                ],
                per_person_total_g=Decimal("20.25"),
                issue_total_raw_g=Decimal("60.75"),
                issue_total_rounded_g=61,
            )
        ],
        source_day_hash="a" * 64,
        generated_by=PydanticObjectId(),
        generated_at=now,
        created_at=now,
        updated_at=now,
    )

    workbook = openpyxl.load_workbook(
        BytesIO(build_menu_requirement_workbook(requirement, school_name="Ліцей №1"))
    )
    sheet = workbook.active

    assert sheet["A1"].value == "МЕНЮ-ВИМОГА"
    assert sheet["B2"].value == "Ліцей №1"
    assert [sheet.cell(row, 1).value for row in range(2, 8)] == [
        "Школа",
        "Вікова група",
        "Дата",
        "Прийом їжі",
        "Меню",
        "Версія",
    ]
    ingredient_row = next(row for row in sheet.iter_rows() if row[0].value == "Морква")
    assert ingredient_row[1].value == 20.25
    assert ingredient_row[-1].value == 61


def test_builds_one_report_sheet_per_school_group() -> None:
    report = MenuRequirementReportResponse.model_validate(
        {
            "school_id": str(PydanticObjectId()),
            "school_name": "Ліцей №1",
            "date_from": "2026-07-01",
            "date_to": "2026-07-31",
            "granularity": "month",
            "meal_type": "lunch",
            "school_group_id": None,
            "status": "complete",
            "missing_dates": [],
            "stale_dates": [],
            "groups": [
                {
                    "school_group_id": str(PydanticObjectId()),
                    "school_group_name": group_name,
                    "age_group": "6-11",
                    "dishes": [
                        {
                            "aggregate_key": "dish:soup",
                            "name": "Овочевий суп",
                            "kind": "dish_card",
                            "recipe_card_number": "12",
                            "yield_amount": "200",
                            "key_reliability": "stable",
                            "children_count_total": 30,
                        }
                    ],
                    "ingredient_rows": [
                        {
                            "key": "ingredient:carrot",
                            "ingredient_id": None,
                            "ingredient_name": "Морква",
                            "cells": [
                                {
                                    "dish_key": "dish:soup",
                                    "net_per_person_g": "20.25",
                                    "issue_total_raw_g": "607.5",
                                    "issue_total_rounded_g": 608,
                                    "breakdown": [],
                                }
                            ],
                            "per_person_total_g": "20.25",
                            "issue_total_raw_g": "607.5",
                            "issue_total_rounded_g": 608,
                        }
                    ],
                }
                for group_name in ("1-А", "1-Б")
            ],
        }
    )

    workbook = openpyxl.load_workbook(BytesIO(build_menu_requirement_report_workbook(report)))

    assert workbook.sheetnames == ["1-А", "1-Б"]
    assert all(sheet["A1"].value == "МЕНЮ-ВИМОГА" for sheet in workbook.worksheets)
    assert [workbook["1-А"].cell(row, 1).value for row in range(2, 7)] == [
        "Школа",
        "Період",
        "Групування",
        "Прийом їжі",
        "Вікова група",
    ]
    first_ingredient_row = next(
        row for row in workbook["1-А"].iter_rows() if row[0].value == "Морква"
    )
    assert first_ingredient_row[1].value == 608
    assert first_ingredient_row[-1].value == 608
