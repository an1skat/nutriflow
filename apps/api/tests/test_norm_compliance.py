from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from beanie import PydanticObjectId

from app.modules.identity.models import AgeGroup, SchoolGroup, User, UserRole
from app.modules.menu_requirements.models import MenuRequirement, MenuRequirementDish
from app.modules.menu_requirements.service import hash_daily_menu
from app.modules.menus.models import (
    DailyMenu,
    DailyMenuItem,
    MealType,
    MenuItemKind,
    MenuItemServingCount,
    MenuPortion,
    Weekday,
)
from app.modules.norm_compliance.registry import get_norm, get_norms
from app.modules.norm_compliance.schemas import (
    ComplianceGroupResponse,
    ComplianceStatus,
    NormComplianceReportResponse,
)
from app.modules.norm_compliance.service import (
    ExpectedRequirement,
    NormComplianceAccessDeniedError,
    NormComplianceValidationError,
    _build_section,
    _get_accessible_school,
    _validate_complete_requirement_week,
    _validate_range,
    calculate_numeric_status,
    contribution_value,
)
from app.modules.norm_compliance.xlsx import build_norm_compliance_workbook
from app.modules.nutrition.domain import (
    NormativeContributionSnapshot,
    NormativeContributionSource,
    NormativeGroupCode,
    NormativeUnit,
)

pytestmark = pytest.mark.no_clean_database


@pytest.mark.parametrize("meal_type", [MealType.BREAKFAST, MealType.LUNCH])
@pytest.mark.parametrize("age_group", list(AgeGroup))
def test_registry_contains_appendix_norms_for_every_supported_age(
    meal_type: MealType,
    age_group: AgeGroup,
) -> None:
    norms = get_norms(meal_type, age_group)

    assert len(norms) == 17
    assert {norm.source_appendix for norm in norms} == {
        "9" if meal_type == MealType.BREAKFAST else "9-1"
    }


def test_registry_preserves_breakfast_and_lunch_values() -> None:
    breakfast_cereal = get_norm(
        MealType.BREAKFAST,
        AgeGroup.SIX_TO_ELEVEN,
        NormativeGroupCode.CEREALS_GRAINS_LEGUMES,
    )
    lunch_meat = get_norm(
        MealType.LUNCH,
        AgeGroup.FOURTEEN_TO_EIGHTEEN,
        NormativeGroupCode.RED_MEAT,
    )

    assert breakfast_cereal is not None
    assert breakfast_cereal.weekly_portions == Decimal("4")
    assert breakfast_cereal.portion_options[0].amount == Decimal("120")
    assert lunch_meat is not None
    assert lunch_meat.weekly_portions == Decimal("3")
    assert lunch_meat.required_amount_per_week == Decimal("360")


@pytest.mark.parametrize(
    ("actual", "expected"),
    [
        (Decimal("99"), ComplianceStatus.COMPLETE),
        (Decimal("100"), ComplianceStatus.COMPLETE),
        (Decimal("101"), ComplianceStatus.COMPLETE),
        (Decimal("98.99"), ComplianceStatus.UNDER),
        (Decimal("101.01"), ComplianceStatus.OVER),
    ],
)
def test_numeric_status_uses_tolerance(actual: Decimal, expected: ComplianceStatus) -> None:
    status, _ = calculate_numeric_status(
        actual,
        Decimal("100"),
        minimum_percent=Decimal("99"),
        maximum_percent=Decimal("101"),
    )
    assert status == expected


@pytest.mark.parametrize(
    "group_code",
    [
        NormativeGroupCode.SALT,
        NormativeGroupCode.SUGAR,
        NormativeGroupCode.COCOA,
        NormativeGroupCode.TEA,
    ],
)
def test_limited_products_are_complete_from_fifty_percent(
    group_code: NormativeGroupCode,
) -> None:
    meal_type = MealType.BREAKFAST
    norm = get_norm(meal_type, AgeGroup.SIX_TO_ELEVEN, group_code)
    assert norm is not None

    status, percent = calculate_numeric_status(
        Decimal("50"),
        Decimal("100"),
        minimum_percent=norm.tolerance.minimum_percent,
        maximum_percent=norm.tolerance.maximum_percent,
    )

    assert percent == Decimal("50")
    assert status == ComplianceStatus.COMPLETE


@pytest.mark.parametrize(
    "group_code",
    [
        NormativeGroupCode.FISH,
        NormativeGroupCode.POULTRY,
        NormativeGroupCode.RED_MEAT,
    ],
)
def test_fish_and_meat_groups_require_exact_weekly_mass(group_code: NormativeGroupCode) -> None:
    meal_type = MealType.LUNCH if group_code == NormativeGroupCode.RED_MEAT else MealType.BREAKFAST
    norm = get_norm(meal_type, AgeGroup.SIX_TO_ELEVEN, group_code)
    assert norm is not None

    status, _ = calculate_numeric_status(
        Decimal("99.99"),
        Decimal("100"),
        minimum_percent=norm.tolerance.minimum_percent,
        maximum_percent=norm.tolerance.maximum_percent,
    )

    assert status == ComplianceStatus.UNDER


def test_regular_groups_allow_a_ten_percent_net_portion_deviation() -> None:
    norm = get_norm(
        MealType.BREAKFAST,
        AgeGroup.SIX_TO_ELEVEN,
        NormativeGroupCode.VEGETABLES,
    )
    assert norm is not None

    assert (
        calculate_numeric_status(
            Decimal("90"),
            Decimal("100"),
            minimum_percent=norm.tolerance.minimum_percent,
            maximum_percent=norm.tolerance.maximum_percent,
        )[0]
        == ComplianceStatus.COMPLETE
    )
    assert (
        calculate_numeric_status(
            Decimal("110"),
            Decimal("100"),
            minimum_percent=norm.tolerance.minimum_percent,
            maximum_percent=norm.tolerance.maximum_percent,
        )[0]
        == ComplianceStatus.COMPLETE
    )


def test_ready_portion_is_compared_with_the_appendix_portion_amount() -> None:
    norm = get_norm(
        MealType.BREAKFAST,
        AgeGroup.SIX_TO_ELEVEN,
        NormativeGroupCode.CEREALS_GRAINS_LEGUMES,
    )
    assert norm is not None
    contribution = NormativeContributionSnapshot(
        group_code=NormativeGroupCode.CEREALS_GRAINS_LEGUMES,
        amount=Decimal("95"),
        unit=NormativeUnit.GRAM,
        portion_equivalent=Decimal("1"),
        source_type=NormativeContributionSource.PORTION_VARIANT,
        source_name="Каша вівсяна",
    )

    assert contribution_value(contribution, norm) == (Decimal("120"), Decimal("1"))


def test_cereal_ingredients_are_counted_once_as_a_dish_portion() -> None:
    group, day, expected = _expected_fixture()
    requirement = _requirement_fixture(
        group,
        day,
        expected.weekly_menu_id,
        amount=None,
    )
    dish = requirement.dishes[0]
    dish.kind = MenuItemKind.DISH_CARD
    dish.name = "Ліниві голубці зі сметаною"
    dish.yield_amount = "180"
    dish.normative_contributions = [
        NormativeContributionSnapshot(
            group_code=NormativeGroupCode.CEREALS_GRAINS_LEGUMES,
            amount=Decimal("30"),
            unit=NormativeUnit.GRAM,
            source_type=NormativeContributionSource.INGREDIENT,
            source_name="Крупа рисова",
        ),
        NormativeContributionSnapshot(
            group_code=NormativeGroupCode.CEREALS_GRAINS_LEGUMES,
            amount=Decimal("20"),
            unit=NormativeUnit.GRAM,
            source_type=NormativeContributionSource.INGREDIENT,
            source_name="Булгур",
        ),
    ]

    section = _build_section(group.age_group, MealType.BREAKFAST, [requirement], [expected])
    cereals = next(
        row
        for row in section.rows
        if row.normative_group_code == NormativeGroupCode.CEREALS_GRAINS_LEGUMES
    )

    assert cereals.actual_amount == Decimal("120")
    assert cereals.actual_portions == Decimal("1")
    assert len(cereals.breakdown) == 1
    assert cereals.breakdown[0].dish_name == dish.name
    assert cereals.breakdown[0].source_name == dish.name
    assert cereals.breakdown[0].source_type == NormativeContributionSource.PORTION_VARIANT


def test_section_reports_complete_row_and_stale_requirement() -> None:
    group, day, expected = _expected_fixture()
    requirement = _requirement_fixture(
        group,
        day,
        expected.weekly_menu_id,
        amount=Decimal("500"),
    )

    complete = _build_section(group.age_group, MealType.BREAKFAST, [requirement], [expected])
    vegetables = next(
        row for row in complete.rows if row.normative_group_code == NormativeGroupCode.VEGETABLES
    )
    assert vegetables.status == ComplianceStatus.COMPLETE
    assert vegetables.actual_amount == Decimal("500")
    assert vegetables.required_amount == Decimal("500")
    assert vegetables.required_portions == Decimal("5")

    requirement.source_day_hash = "0" * 64
    stale = _build_section(group.age_group, MealType.BREAKFAST, [requirement], [expected])
    stale_vegetables = next(
        row for row in stale.rows if row.normative_group_code == NormativeGroupCode.VEGETABLES
    )
    assert stale.stale_dates == [date(2026, 7, 6)]
    assert stale_vegetables.status == ComplianceStatus.STALE


def test_weekly_norms_are_not_prorated_by_available_days() -> None:
    group, day, expected = _expected_fixture()
    requirement = _requirement_fixture(
        group,
        day,
        expected.weekly_menu_id,
        amount=Decimal("100"),
    )

    section = _build_section(group.age_group, MealType.BREAKFAST, [requirement], [expected])
    juice = next(
        row for row in section.rows if row.normative_group_code == NormativeGroupCode.JUICES
    )

    assert juice.required_portions == Decimal("1")
    assert juice.required_amount == Decimal("200")


def test_section_reports_missing_and_unmapped_requirements() -> None:
    group, day, expected = _expected_fixture()
    missing = _build_section(group.age_group, MealType.BREAKFAST, [], [expected])
    assert missing.missing_dates == [date(2026, 7, 6)]
    assert all(row.status == ComplianceStatus.MISSING for row in missing.rows)

    requirement = _requirement_fixture(group, day, expected.weekly_menu_id, amount=None)
    unmapped = _build_section(
        group.age_group,
        MealType.BREAKFAST,
        [requirement],
        [expected],
    )
    assert unmapped.unmapped_items[0].item_name == "Овочевий салат"
    assert all(row.status == ComplianceStatus.UNMAPPED for row in unmapped.rows)


@pytest.mark.asyncio
async def test_school_user_cannot_read_norm_compliance() -> None:
    current_user = User.model_construct(
        id=PydanticObjectId(),
        role=UserRole.SCHOOL_USER,
        school_id=PydanticObjectId(),
    )

    with pytest.raises(NormComplianceAccessDeniedError, match="Norm compliance access denied"):
        await _get_accessible_school(current_user, current_user.school_id)


def test_report_range_must_be_one_monday_to_friday_week() -> None:
    _validate_range(date(2026, 7, 6), date(2026, 7, 10))

    with pytest.raises(NormComplianceValidationError, match="Monday-to-Friday"):
        _validate_range(date(2026, 7, 7), date(2026, 7, 11))

    with pytest.raises(NormComplianceValidationError, match="Monday-to-Friday"):
        _validate_range(date(2026, 7, 6), date(2026, 7, 9))


def test_report_requires_menu_requirements_for_all_five_weekdays() -> None:
    group, base_day, _expected = _expected_fixture()
    weekly_menu_id = PydanticObjectId()
    expected: list[ExpectedRequirement] = []
    requirements: list[MenuRequirement] = []

    for offset, weekday in enumerate(list(Weekday)[:5]):
        service_date = date(2026, 7, 6) + timedelta(days=offset)
        day = base_day.model_copy(
            update={"weekday": weekday, "date": service_date},
            deep=True,
        )
        expected_item = ExpectedRequirement(
            weekly_menu_id=weekly_menu_id,
            weekday=weekday,
            service_date=service_date,
            meal_type=MealType.BREAKFAST,
            school_group=group,
            day=day,
        )
        expected.append(expected_item)
        requirements.append(
            _requirement_fixture(
                group,
                day,
                weekly_menu_id,
                amount=Decimal("500"),
            )
        )

    _validate_complete_requirement_week(
        requirements,
        expected,
        date(2026, 7, 6),
        date(2026, 7, 10),
    )

    with pytest.raises(NormComplianceValidationError, match="all five weekdays"):
        _validate_complete_requirement_week(
            requirements[:4],
            expected,
            date(2026, 7, 6),
            date(2026, 7, 10),
        )

    with pytest.raises(NormComplianceValidationError, match="all five weekdays"):
        _validate_complete_requirement_week(
            requirements[:4],
            expected[:4],
            date(2026, 7, 6),
            date(2026, 7, 10),
        )


def test_builds_norm_compliance_workbook() -> None:
    group, day, expected = _expected_fixture()
    requirement = _requirement_fixture(
        group,
        day,
        expected.weekly_menu_id,
        amount=Decimal("500"),
    )
    section = _build_section(group.age_group, MealType.BREAKFAST, [requirement], [expected])
    report = NormComplianceReportResponse(
        school_id=PydanticObjectId(),
        school_name="Ліцей №1",
        date_from=date(2026, 7, 6),
        date_to=date(2026, 7, 10),
        status=section.status,
        groups=[
            ComplianceGroupResponse(
                school_group_id=group.id,
                school_group_name=group.name,
                age_group=group.age_group,
                status=section.status,
                sections=[section],
            )
        ],
    )

    workbook = openpyxl.load_workbook(BytesIO(build_norm_compliance_workbook(report)))
    sheet = workbook.active

    assert sheet["A1"].value == "ДОТРИМАННЯ НОРМ ХАРЧУВАННЯ"
    assert sheet["B2"].value == "Ліцей №1"
    vegetables = next(row for row in sheet.iter_rows() if row[0].value == "Овочі")
    assert vegetables[4].value == 500
    assert vegetables[5].value == 500
    assert vegetables[-1].value == "В нормі"


def _expected_fixture() -> tuple[SchoolGroup, DailyMenu, ExpectedRequirement]:
    group = SchoolGroup(name="6-11", age_group=AgeGroup.SIX_TO_ELEVEN)
    menu_item = DailyMenuItem(
        position=1,
        kind=MenuItemKind.PRODUCT,
        name="Овочевий салат",
        product_name_snapshot="Овочевий салат",
        portions=[MenuPortion(age_group=group.age_group, yield_amount="100")],
        servings=[
            MenuItemServingCount(
                school_group_id=group.id,
                age_group=group.age_group,
                children_count=20,
            )
        ],
    )
    day = DailyMenu(weekday=Weekday.MONDAY, date=date(2026, 7, 6), items=[menu_item])
    expected = ExpectedRequirement(
        weekly_menu_id=PydanticObjectId(),
        weekday=day.weekday,
        service_date=day.date,
        meal_type=MealType.BREAKFAST,
        school_group=group,
        day=day,
    )
    return group, day, expected


def _requirement_fixture(
    group: SchoolGroup,
    day: DailyMenu,
    weekly_menu_id: PydanticObjectId,
    *,
    amount: Decimal | None,
) -> MenuRequirement:
    item = day.items[0]
    contributions = (
        [
            NormativeContributionSnapshot(
                group_code=NormativeGroupCode.VEGETABLES,
                amount=amount,
                unit=NormativeUnit.GRAM,
                source_type=NormativeContributionSource.PRODUCT,
                source_name="Овочевий салат",
            )
        ]
        if amount is not None
        else []
    )
    return MenuRequirement.model_construct(
        id=PydanticObjectId(),
        school_id=PydanticObjectId(),
        weekly_menu_id=weekly_menu_id,
        menu_title="Тижневе меню",
        meal_type=MealType.BREAKFAST,
        weekday=day.weekday,
        service_date=day.date,
        school_group_id=group.id,
        school_group_name=group.name,
        age_group=group.age_group,
        dishes=[
            MenuRequirementDish(
                menu_item_id=item.id,
                position=1,
                kind=MenuItemKind.PRODUCT,
                name=item.name,
                yield_amount="100",
                children_count=20,
                normative_contributions=contributions,
            )
        ],
        source_day_hash=hash_daily_menu(day),
        generated_by=PydanticObjectId(),
    )


@pytest.mark.parametrize(
    "meal_type, entries, actual, status",
    [
        (MealType.LUNCH, [("milk", "18", "ml"), ("milk", "4", "ml")], "0", ComplianceStatus.UNDER),
        (
            MealType.LUNCH,
            [
                ("sour_cream", "4", "g"),
                ("hard_cheese", "15", "g"),
                ("milk", "48", "ml"),
                ("milk", "6", "ml"),
            ],
            "1.16",
            ComplianceStatus.UNDER,
        ),
        (MealType.LUNCH, [("sour_cream", "14", "g")], "0.56", ComplianceStatus.UNDER),
        (MealType.LUNCH, [("soft_cheese", "75", "g")], "1", ComplianceStatus.UNDER),
        (MealType.BREAKFAST, [("milk", "200", "ml")], "1", ComplianceStatus.UNDER),
        (MealType.LUNCH, [("unknown", "18", "ml")], "0", ComplianceStatus.UNMAPPED),
        (MealType.LUNCH, [(None, "18", "ml")], "0", ComplianceStatus.UNMAPPED),
        (MealType.LUNCH, [("milk", "18", "g")], "0", ComplianceStatus.UNMAPPED),
        (MealType.LUNCH, [("hard_cheese", "15", "ml")], "0", ComplianceStatus.UNMAPPED),
    ],
)
def test_dairy_variant_applicability(meal_type, entries, actual, status) -> None:
    group, day, expected = _expected_fixture()
    requirement = _requirement_fixture(group, day, expected.weekly_menu_id, amount=None)
    requirement.meal_type = meal_type
    requirement.dishes[0].normative_contributions = [
        NormativeContributionSnapshot(
            group_code=NormativeGroupCode.DAIRY,
            product_variant=variant,
            amount=Decimal(amount),
            unit=unit,
            source_type=NormativeContributionSource.INGREDIENT,
            source_name=variant or "Missing variant",
        )
        for variant, amount, unit in entries
    ]
    section = _build_section(group.age_group, meal_type, [requirement], [])
    dairy = next(
        row for row in section.rows if row.normative_group_code == NormativeGroupCode.DAIRY
    )
    assert dairy.actual_portions == Decimal(actual)
    assert dairy.status == status
    assert len(dairy.breakdown) == (
        len(entries)
        if meal_type == MealType.BREAKFAST
        else sum(
            variant in {"soft_cheese", "hard_cheese", "sour_cream"} and unit == "g"
            for variant, _, unit in entries
        )
    )


def test_lunch_dairy_production_week_ignores_milk() -> None:
    group, base_day, expected = _expected_fixture()
    requirements = []
    daily_entries = [
        [("milk", "18", "ml"), ("milk", "4", "ml")],
        [
            ("sour_cream", "4", "g"),
            ("hard_cheese", "15", "g"),
            ("milk", "48", "ml"),
            ("milk", "6", "ml"),
        ],
        [],
        [],
        [("sour_cream", "14", "g")],
    ]
    for offset, entries in enumerate(daily_entries):
        day = base_day.model_copy(
            update={
                "date": date(2026, 9, 14) + timedelta(days=offset),
                "weekday": list(Weekday)[offset],
            }
        )
        requirement = _requirement_fixture(group, day, expected.weekly_menu_id, amount=None)
        requirement.school_id = PydanticObjectId("6a848408f9867d3ec9e6c866")
        requirement.meal_type = MealType.LUNCH
        requirement.dishes[0].normative_contributions = [
            NormativeContributionSnapshot(
                group_code=NormativeGroupCode.DAIRY,
                product_variant=variant,
                amount=Decimal(amount),
                unit=unit,
                source_type=NormativeContributionSource.INGREDIENT,
                source_name=variant,
            )
            for variant, amount, unit in entries
        ]
        if not entries:
            requirement.dishes = []
        requirements.append(requirement)
    section = _build_section(group.age_group, MealType.LUNCH, requirements, [])
    dairy = next(
        row for row in section.rows if row.normative_group_code == NormativeGroupCode.DAIRY
    )
    assert dairy.actual_portions == dairy.actual_amount == Decimal("1.72")
    assert dairy.required_portions == dairy.required_amount == Decimal("3")
    assert dairy.status == ComplianceStatus.UNDER
    assert [(item.amount, item.service_date) for item in dairy.breakdown] == [
        (Decimal("4"), date(2026, 9, 15)),
        (Decimal("15"), date(2026, 9, 15)),
        (Decimal("14"), date(2026, 9, 18)),
    ]


@pytest.mark.parametrize("week_count", [1, 2, 3, 4, 5])
def test_selected_period_scales_mass_portions_and_percentage(week_count: int) -> None:
    group, day, expected = _expected_fixture()
    requirement = _requirement_fixture(group, day, expected.weekly_menu_id, amount=Decimal("100"))
    baseline = _build_section(group.age_group, MealType.BREAKFAST, [requirement], [expected])
    scaled = _build_section(
        group.age_group,
        MealType.BREAKFAST,
        [requirement] * week_count,
        [expected],
        week_count=week_count,
    )
    for base, row in zip(baseline.rows, scaled.rows, strict=True):
        assert row.required_amount == base.required_amount * week_count
        assert row.required_portions == base.required_portions * week_count
        assert row.actual_amount == base.actual_amount * week_count
        assert row.actual_portions == base.actual_portions * week_count
        assert row.percent == base.percent
        assert row.deviation == base.deviation * week_count
        assert row.status == base.status


@pytest.mark.parametrize(
    "dates",
    [
        [],
        [date(2026, 7, 1)],
        [date(2026, 7, 1)] * 5,
        [date(2026, 7, n) for n in [1, 2, 3, 4, 6]],
        [date(2026, 7, n) for n in [1, 2, 3, 6]] + [date(2026, 8, 3)],
    ],
)
def test_selected_dates_reject_invalid_counts_duplicates_weekends_and_months(dates) -> None:
    from app.modules.norm_compliance.service import _validate_selected_dates

    with pytest.raises(NormComplianceValidationError):
        _validate_selected_dates(
            dates,
            min(dates) if dates else date(2026, 7, 1),
            max(dates) if dates else date(2026, 7, 31),
        )


def test_selected_dates_allow_gaps_and_reject_inconsistent_bounds() -> None:
    from app.modules.norm_compliance.service import _validate_selected_dates

    dates = [date(2026, 7, n) for n in [1, 3, 7, 9, 15]]
    _validate_selected_dates(dates, dates[0], dates[-1])
    with pytest.raises(NormComplianceValidationError):
        _validate_selected_dates(dates, dates[0], date(2026, 7, 31))


@pytest.mark.asyncio
async def test_report_filters_unselected_days_and_rejects_not_served(monkeypatch) -> None:
    from dataclasses import replace

    from app.modules.identity.models import School
    from app.modules.norm_compliance import service

    group, day, expected = _expected_fixture()
    dates = [date(2026, 7, n) for n in [1, 3, 6, 8, 10, 13, 15, 17, 22, 24]]
    expected_items = [
        replace(expected, service_date=value, day=day.model_copy(update={"date": value}))
        for value in dates + [date(2026, 7, 7)]
    ]
    requirements = [
        _requirement_fixture(group, item.day, item.weekly_menu_id, amount=Decimal("100"))
        for item in expected_items
    ]
    school = School.model_construct(id=PydanticObjectId(), name="School", groups=[group])

    async def get_school(*args, **kwargs):
        return school

    async def get_requirements(*args, **kwargs):
        return requirements

    async def get_expected(*args, **kwargs):
        return expected_items

    async def apply_rules(*args):
        pass

    monkeypatch.setattr(service, "_get_accessible_school", get_school)
    monkeypatch.setattr(service, "_find_requirements", get_requirements)
    monkeypatch.setattr(service, "_find_expected_requirements", get_expected)
    monkeypatch.setattr(service, "_apply_manual_ingredient_rules", apply_rules)
    owner = User.model_construct(role=UserRole.OWNER)
    report = await service.get_norm_compliance_report(
        school.id, dates[0], dates[-1], owner, dates=dates
    )
    assert report.selected_dates == dates
    section = report.groups[0].sections[0]
    vegetables = next(
        row for row in section.rows if row.normative_group_code == NormativeGroupCode.VEGETABLES
    )
    assert vegetables.actual_amount == Decimal("1000")
    assert {entry.service_date for entry in vegetables.breakdown} == set(dates)
    norm = get_norm(MealType.BREAKFAST, group.age_group, NormativeGroupCode.VEGETABLES)
    assert vegetables.required_amount == norm.required_amount_per_week * 2
    expected_items[0] = replace(expected_items[0], day=day.model_copy(update={"not_served": True}))
    with pytest.raises(NormComplianceValidationError, match="не харчувалися"):
        await service.get_norm_compliance_report(school.id, dates[0], dates[-1], owner, dates=dates)


@pytest.mark.parametrize("path", ["/report", "/report/export.xlsx"])
def test_report_routes_preserve_selected_dates(path: str, monkeypatch: pytest.MonkeyPatch) -> None:
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from app.modules.auth.dependencies import get_current_user
    from app.modules.norm_compliance import router as routes

    dates = [date(2026, 7, n) for n in [1, 3, 6, 8, 10]]
    school_id = PydanticObjectId()
    received = []

    async def get_report(*args, **kwargs):
        received.append(kwargs["dates"])
        return NormComplianceReportResponse(
            school_id=school_id,
            school_name="School",
            date_from=dates[0],
            date_to=dates[-1],
            selected_dates=kwargs["dates"],
            status=ComplianceStatus.COMPLETE,
            groups=[],
        )

    monkeypatch.setattr(routes, "get_norm_compliance_report", get_report)
    app = FastAPI()
    app.include_router(routes.router)
    app.dependency_overrides[get_current_user] = lambda: User.model_construct(role=UserRole.OWNER)
    client = TestClient(app)
    params = [
        ("school_id", str(school_id)),
        ("date_from", dates[0].isoformat()),
        ("date_to", dates[-1].isoformat()),
    ]
    params.extend(("dates", day.isoformat()) for day in dates)
    response = client.get(path, params=params)
    assert response.status_code == 200
    assert received == [dates]
    if path.endswith("xlsx"):
        workbook = openpyxl.load_workbook(BytesIO(response.content))
        assert workbook.active["B3"].value == "Обрані дні: " + ", ".join(
            day.strftime("%d.%m.%Y") for day in dates
        )
    else:
        assert response.json()["selected_dates"] == [day.isoformat() for day in dates]
    invalid_params = [(key, value) for key, value in params if key != "dates"]
    assert client.get(path, params=invalid_params + [("dates", "bad-date")]).status_code == 422
    assert len(received) == 1


def test_completeness_checks_selected_days_instead_of_gaps() -> None:
    from dataclasses import replace

    group, base_day, expected = _expected_fixture()
    dates = [date(2026, 7, n) for n in [1, 3, 6, 8, 10]]
    items = [
        replace(expected, service_date=value, day=base_day.model_copy(update={"date": value}))
        for value in dates
    ]
    requirements = [
        _requirement_fixture(group, item.day, item.weekly_menu_id, amount=None) for item in items
    ]
    _validate_complete_requirement_week(
        requirements, items, dates[0], dates[-1], dates=dates, meal_type=MealType.BREAKFAST
    )
    with pytest.raises(
        NormComplianceValidationError, match="10.07.2026 — Сніданок: немає меню-вимоги"
    ):
        _validate_complete_requirement_week(
            requirements, items[:-1], dates[0], dates[-1], dates=dates
        )


@pytest.mark.parametrize(
    "dates",
    [
        [date(2026, 9, n) for n in [28, 29, 30]] + [date(2026, 10, n) for n in [1, 2]],
        [date(2026, 12, n) for n in [28, 29, 30, 31]] + [date(2027, 1, 1)],
        [date(2026, 6, n) for n in [29, 30]] + [date(2026, 7, n) for n in [1, 2, 3]],
    ],
)
def test_selected_dates_allow_complete_calendar_weeks_across_months(dates) -> None:
    from app.modules.norm_compliance.service import _validate_selected_dates

    _validate_selected_dates(dates, dates[0], dates[-1])
