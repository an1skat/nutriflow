from copy import deepcopy
from datetime import UTC, datetime
from datetime import date as Date
from io import BytesIO

import openpyxl
import pytest
from beanie import PydanticObjectId
from fastapi.testclient import TestClient
from pymongo import MongoClient

from app.core.config import get_settings
from app.modules.menu_requirements.service import _menu_day_service_date, resolve_service_date
from app.modules.menu_requirements.utils import hash_daily_menu
from app.modules.menus.day_closure import (
    _current_school_workweek,
    _resolve_auto_close_service_date,
)
from app.modules.menus.models import DailyMenu, MealType, Weekday, WeeklyMenu


def login(client: TestClient, identifier: str, password: str) -> None:
    response = client.post(
        "/api/v1/auth/login",
        json={"identifier": identifier, "password": password},
    )
    assert response.status_code == 204


def csrf_headers(client: TestClient) -> dict[str, str]:
    token = client.cookies.get(get_settings().csrf_cookie_name)
    assert token is not None
    return {"X-CSRF-Token": token}


@pytest.mark.no_clean_database
def test_service_date_uses_persisted_day_date_without_weekday_shift() -> None:
    day = DailyMenu.model_construct(
        weekday=Weekday.MONDAY,
        date=Date(2026, 7, 7),
        items=[],
    )
    menu = WeeklyMenu.model_construct(
        title="Меню з явною датою дня",
        meal_type=MealType.LUNCH,
        starts_on=Date(2026, 7, 6),
        days=[day],
    )

    assert resolve_service_date(menu, day, Date(2026, 7, 6)) == Date(2026, 7, 7)
    assert _menu_day_service_date(menu, day) == Date(2026, 7, 7)


@pytest.mark.no_clean_database
def test_service_date_derives_from_week_start_when_day_date_is_missing() -> None:
    day = DailyMenu.model_construct(
        weekday=Weekday.WEDNESDAY,
        date=None,
        items=[],
    )
    menu = WeeklyMenu.model_construct(
        title="Меню з датою початку тижня",
        meal_type=MealType.LUNCH,
        starts_on=Date(2026, 7, 6),
        days=[day],
    )

    assert resolve_service_date(menu, day, Date(2026, 7, 6)) == Date(2026, 7, 8)
    assert _menu_day_service_date(menu, day) == Date(2026, 7, 8)


@pytest.mark.no_clean_database
def test_current_school_workweek_is_monday_through_friday() -> None:
    assert _current_school_workweek(Date(2026, 8, 1)) == (
        Date(2026, 7, 27),
        Date(2026, 7, 31),
    )


@pytest.mark.no_clean_database
def test_reopened_day_is_not_auto_closed_again() -> None:
    day = DailyMenu.model_construct(
        weekday=Weekday.MONDAY,
        date=Date(2026, 7, 6),
        reopened_at=datetime(2026, 7, 20, 10, 0, tzinfo=UTC),
        items=[],
    )
    menu = WeeklyMenu.model_construct(
        title="Повторно відкрите меню",
        meal_type=MealType.LUNCH,
        starts_on=Date(2026, 7, 6),
        days=[day],
    )

    assert _resolve_auto_close_service_date(menu, day, Date(2026, 7, 20)) is None


@pytest.mark.no_clean_database
def test_requirement_hash_ignores_close_notification_delivery_state() -> None:
    day = DailyMenu.model_construct(
        weekday=Weekday.MONDAY,
        date=Date(2026, 7, 6),
        items=[],
        close_notification_pending=False,
        close_notification_sent_at=None,
    )
    initial_hash = hash_daily_menu(day)

    day.close_notification_pending = True
    day.close_notification_sent_at = datetime(2026, 7, 6, 16, 0, tzinfo=UTC)
    day.reopened_at = datetime(2026, 7, 6, 17, 0, tzinfo=UTC)
    day.reopened_by = PydanticObjectId()

    assert hash_daily_menu(day) == initial_hash


def create_confirmed_dish(
    client: TestClient,
    *,
    ingredient_id: str,
) -> tuple[str, str]:
    card_response = client.post(
        "/api/v1/recipes/dish-cards",
        json={"card_number": "REQ-1", "name": "Овочевий суп"},
        headers=csrf_headers(client),
    )
    assert card_response.status_code == 201
    card_id = card_response.json()["id"]
    variant_id = PydanticObjectId()

    version_response = client.post(
        f"/api/v1/recipes/dish-cards/{card_id}/versions",
        json={
            "portion_variants": [
                {
                    "id": str(variant_id),
                    "age_group": "6-11",
                    "output_grams": "200",
                    "nutrition": {},
                }
            ],
            "ingredient_amounts": [
                {
                    "ingredient_id": ingredient_id,
                    "ingredient_name_snapshot": "Морква",
                    "gross_amount": "25",
                    "net_amount": "20.25",
                    "unit": "g",
                    "portion_variant_id": str(variant_id),
                }
            ],
        },
        headers=csrf_headers(client),
    )
    assert version_response.status_code == 201
    version_id = version_response.json()["id"]
    confirm_response = client.post(
        f"/api/v1/recipes/dish-card-versions/{version_id}/confirm",
        headers=csrf_headers(client),
    )
    assert confirm_response.status_code == 200
    return card_id, str(variant_id)


def publish_school_menu(
    client: TestClient,
    *,
    school_id: str,
    card_id: str,
    variant_id: str,
) -> dict:
    create_response = client.post(
        "/api/v1/menus/weekly",
        json={
            "title": "Меню для вимоги",
            "meal_type": "lunch",
            "starts_on": "2026-07-06",
            "ends_on": "2026-07-10",
            "days": [
                {
                    "weekday": "monday",
                    "date": "2026-07-06",
                    "items": [
                        {
                            "position": 1,
                            "kind": "dish_card",
                            "recipe_card_number": "REQ-1",
                            "dish_card_id": card_id,
                            "name": "Овочевий суп",
                            "allergen_codes": [],
                            "portions": [
                                {
                                    "age_group": "6-11",
                                    "yield_amount": "200",
                                    "dish_card_portion_variant_id": variant_id,
                                    "nutrition": {},
                                }
                            ],
                        }
                    ],
                }
            ],
        },
        headers=csrf_headers(client),
    )
    assert create_response.status_code == 201
    source_id = create_response.json()["id"]

    publish_response = client.post(
        f"/api/v1/menus/weekly/{source_id}/publish",
        json={"school_ids": [school_id]},
        headers=csrf_headers(client),
    )
    assert publish_response.status_code == 200
    school_menu_id = publish_response.json()["created_menu_ids"][0]

    menu_response = client.get(f"/api/v1/menus/weekly/{school_menu_id}")
    assert menu_response.status_code == 200
    return menu_response.json()


def prepare_school_menu_with_count(
    client: TestClient,
    *,
    menu: dict,
    group_id: str,
    count: int,
) -> dict:
    days = menu["days"]
    days[0]["items"][0]["servings"] = [
        {
            "school_group_id": group_id,
            "age_group": "6-11",
            "children_count": count,
        }
    ]
    response = client.patch(
        f"/api/v1/menus/weekly/{menu['id']}",
        json={"days": days, "revision": menu["revision"]},
        headers=csrf_headers(client),
    )
    assert response.status_code == 200
    return response.json()


def test_school_generates_and_regenerates_menu_requirement(seeded_client) -> None:
    client, identities = seeded_client
    login(client, identities.admin.username, identities.admin_password)

    ingredient_response = client.post(
        "/api/v1/recipes/ingredients",
        json={"name": "Морква", "unit": "g"},
        headers=csrf_headers(client),
    )
    assert ingredient_response.status_code == 201
    ingredient_id = ingredient_response.json()["id"]
    card_id, variant_id = create_confirmed_dish(
        client,
        ingredient_id=ingredient_id,
    )
    menu = publish_school_menu(
        client,
        school_id=str(identities.own_school.id),
        card_id=card_id,
        variant_id=variant_id,
    )

    login(client, identities.school_user.username, identities.school_user_password)
    group_id = str(identities.own_school.groups[0].id)
    menu = prepare_school_menu_with_count(
        client,
        menu=menu,
        group_id=group_id,
        count=3,
    )

    generate_response = client.post(
        "/api/v1/menu-requirements/generate",
        json={
            "weekly_menu_id": menu["id"],
            "weekday": "monday",
            "service_date": "2026-07-06",
        },
        headers=csrf_headers(client),
    )
    assert generate_response.status_code == 200
    requirement = generate_response.json()["items"][0]
    assert requirement["school_group_id"] == group_id
    assert requirement["service_date"] == "2026-07-06"
    assert requirement["dishes"][0]["children_count"] == 3
    carrot = next(
        row for row in requirement["ingredient_rows"] if row["ingredient_id"] == ingredient_id
    )
    assert carrot["cells"][0]["net_per_person_g"] == "20.25"
    assert carrot["cells"][0]["gross_per_person_g"] == "25"
    assert carrot["per_person_total_g"] == "20.25"
    assert carrot["issue_total_raw_g"] == "60.75"
    assert carrot["issue_total_rounded_g"] == 61
    assert carrot["gross_per_person_total_g"] == "25"
    assert carrot["gross_issue_total_raw_g"] == "75"
    assert carrot["gross_issue_total_rounded_g"] == 75

    school_export_response = client.get(
        f"/api/v1/menu-requirements/{requirement['id']}/export.xlsx"
    )
    assert school_export_response.status_code == 200
    assert school_export_response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    exported_workbook = openpyxl.load_workbook(BytesIO(school_export_response.content))
    exported_sheet = exported_workbook.active
    assert exported_sheet["A1"].value == "МЕНЮ-ВИМОГА"
    assert exported_sheet["B2"].value == identities.own_school.name
    assert any(cell.value == "Морква" for row in exported_sheet.iter_rows() for cell in row)

    gross_export_response = client.get(
        f"/api/v1/menu-requirements/{requirement['id']}/export.xlsx",
        params={"amount_basis": "gross"},
    )
    assert gross_export_response.status_code == 200
    gross_sheet = openpyxl.load_workbook(BytesIO(gross_export_response.content)).active
    gross_carrot_row = next(row for row in gross_sheet.iter_rows() if row[0].value == "Морква")
    assert gross_carrot_row[1].value == 25
    assert gross_carrot_row[-1].value == 75

    regenerate_response = client.post(
        "/api/v1/menu-requirements/generate",
        json={
            "weekly_menu_id": menu["id"],
            "weekday": "monday",
            "service_date": "2026-07-06",
        },
        headers=csrf_headers(client),
    )
    assert regenerate_response.status_code == 200
    assert regenerate_response.json()["items"][0]["revision"] == 2

    login(client, identities.lower_admin.username, identities.lower_admin_password)
    lower_admin_export_response = client.get(
        f"/api/v1/menu-requirements/{requirement['id']}/export.xlsx"
    )
    assert lower_admin_export_response.status_code == 200

    update_payload = {
        "ingredient_rows": [
            {
                "key": carrot["key"],
                "ingredient_name": "Морква очищена",
                "cells": [
                    {
                        "menu_item_id": requirement["dishes"][0]["menu_item_id"],
                        "net_per_person_g": "21.5",
                    }
                ],
            }
        ]
    }
    school_update_response = client.patch(
        f"/api/v1/menu-requirements/{requirement['id']}",
        json=update_payload,
        headers=csrf_headers(client),
    )
    assert school_update_response.status_code == 403

    login(client, identities.admin.username, identities.admin_password)
    owner_update_response = client.patch(
        f"/api/v1/menu-requirements/{requirement['id']}",
        json=update_payload,
        headers=csrf_headers(client),
    )
    assert owner_update_response.status_code == 200
    updated_requirement = owner_update_response.json()
    updated_carrot = updated_requirement["ingredient_rows"][0]
    assert updated_requirement["revision"] == 3
    assert updated_carrot["ingredient_name"] == "Морква очищена"
    assert updated_carrot["per_person_total_g"] == "21.5"
    assert updated_carrot["issue_total_raw_g"] == "64.5"
    assert updated_carrot["issue_total_rounded_g"] == 65
    assert updated_carrot["gross_per_person_total_g"] == "25"
    assert updated_carrot["gross_issue_total_raw_g"] == "75"
    assert updated_carrot["gross_issue_total_rounded_g"] == 75

    owner_export_response = client.get(f"/api/v1/menu-requirements/{requirement['id']}/export.xlsx")
    assert owner_export_response.status_code == 200

    list_response = client.get(f"/api/v1/menu-requirements?weekly_menu_id={menu['id']}")
    assert list_response.status_code == 200
    assert list_response.json()["total"] == 1

    login(client, identities.school_user.username, identities.school_user_password)
    school_delete_response = client.delete(
        f"/api/v1/menu-requirements/{requirement['id']}",
        headers=csrf_headers(client),
    )
    assert school_delete_response.status_code == 403

    login(client, identities.admin.username, identities.admin_password)
    owner_delete_response = client.delete(
        f"/api/v1/menu-requirements/{requirement['id']}",
        headers=csrf_headers(client),
    )
    assert owner_delete_response.status_code == 204

    missing_response = client.get(f"/api/v1/menu-requirements/{requirement['id']}")
    assert missing_response.status_code == 404

    empty_list_response = client.get(f"/api/v1/menu-requirements?weekly_menu_id={menu['id']}")
    assert empty_list_response.status_code == 200
    assert empty_list_response.json()["total"] == 0

    create_technologist_response = client.post(
        "/api/v1/admin/admins",
        json={
            "username": "tech.user",
            "email": "tech.user@example.com",
            "password": "tech-password-123",
            "role": "TECHNOLOGIST",
            "permissions": [],
        },
        headers=csrf_headers(client),
    )
    assert create_technologist_response.status_code == 201

    login(client, identities.school_user.username, identities.school_user_password)
    second_generate_response = client.post(
        "/api/v1/menu-requirements/generate",
        json={
            "weekly_menu_id": menu["id"],
            "weekday": "monday",
            "service_date": "2026-07-06",
        },
        headers=csrf_headers(client),
    )
    assert second_generate_response.status_code == 200
    second_requirement = second_generate_response.json()["items"][0]

    login(client, "tech.user", "tech-password-123")
    technologist_export_response = client.get(
        f"/api/v1/menu-requirements/{second_requirement['id']}/export.xlsx"
    )
    assert technologist_export_response.status_code == 200
    technologist_delete_response = client.delete(
        f"/api/v1/menu-requirements/{second_requirement['id']}",
        headers=csrf_headers(client),
    )
    assert technologist_delete_response.status_code == 204


def test_school_closes_day_and_admin_reopens_it(
    seeded_client,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    client, identities = seeded_client
    login(client, identities.admin.username, identities.admin_password)

    ingredient_response = client.post(
        "/api/v1/recipes/ingredients",
        json={"name": "Морква", "unit": "g"},
        headers=csrf_headers(client),
    )
    assert ingredient_response.status_code == 201
    ingredient_id = ingredient_response.json()["id"]
    card_id, variant_id = create_confirmed_dish(
        client,
        ingredient_id=ingredient_id,
    )
    menu = publish_school_menu(
        client,
        school_id=str(identities.own_school.id),
        card_id=card_id,
        variant_id=variant_id,
    )

    login(client, identities.school_user.username, identities.school_user_password)
    group_id = str(identities.own_school.groups[0].id)
    menu = prepare_school_menu_with_count(
        client,
        menu=menu,
        group_id=group_id,
        count=5,
    )

    close_response = client.post(
        f"/api/v1/menus/weekly/{menu['id']}/days/monday/close",
        headers=csrf_headers(client),
    )
    assert close_response.status_code == 200
    closed_menu = close_response.json()
    closed_day = closed_menu["days"][0]
    assert closed_day["closed_at"] is not None
    assert closed_day["closed_by"] == str(identities.school_user.id)
    assert closed_day["close_reason"] == "manual"
    assert closed_menu["revision"] == menu["revision"] + 1

    duplicate_close_response = client.post(
        f"/api/v1/menus/weekly/{menu['id']}/days/monday/close",
        headers=csrf_headers(client),
    )
    assert duplicate_close_response.status_code == 200
    duplicate_closed_menu = duplicate_close_response.json()
    assert duplicate_closed_menu["days"][0]["closed_at"] == closed_day["closed_at"]
    assert duplicate_closed_menu["revision"] == closed_menu["revision"]

    settings = get_settings()
    mongo_client = MongoClient(settings.mongo_uri, tz_aware=True)
    try:
        stored_menu = mongo_client[settings.mongo_db]["weekly_menus"].find_one(
            {"_id": PydanticObjectId(menu["id"])}
        )
    finally:
        mongo_client.close()
    assert stored_menu is not None
    assert stored_menu["days"][0]["close_notification_pending"] is True
    assert stored_menu["days"][0]["close_notification_sent_at"] is None

    requirements_response = client.get(f"/api/v1/menu-requirements?weekly_menu_id={menu['id']}")
    assert requirements_response.status_code == 200
    requirements = requirements_response.json()["items"]
    assert len(requirements) == 1
    assert requirements[0]["dishes"][0]["children_count"] == 5

    changed_days = closed_menu["days"]
    changed_days[0]["items"][0]["servings"][0]["children_count"] = 6
    edit_response = client.patch(
        f"/api/v1/menus/weekly/{menu['id']}",
        json={"days": changed_days, "revision": closed_menu["revision"]},
        headers=csrf_headers(client),
    )
    assert edit_response.status_code == 400
    assert edit_response.json()["detail"] == "Closed daily menus cannot be changed"

    added_days = deepcopy(closed_menu["days"])
    added_days[0]["items"].append(
        {
            "position": 2,
            "kind": "product",
            "product_name_snapshot": "Хліб",
            "name": "Хліб",
            "allergen_codes": [],
            "portions": [
                {
                    "age_group": "6-11",
                    "yield_amount": "30",
                    "nutrition": {},
                }
            ],
            "servings": [
                {
                    "school_group_id": group_id,
                    "age_group": "6-11",
                    "children_count": 5,
                }
            ],
        }
    )
    add_to_closed_day_response = client.patch(
        f"/api/v1/menus/weekly/{menu['id']}",
        json={"days": added_days, "revision": closed_menu["revision"]},
        headers=csrf_headers(client),
    )
    assert add_to_closed_day_response.status_code == 400
    assert add_to_closed_day_response.json()["detail"] == "Closed daily menus cannot be changed"

    regenerate_response = client.post(
        "/api/v1/menu-requirements/generate",
        json={
            "weekly_menu_id": menu["id"],
            "weekday": "monday",
            "service_date": "2026-07-06",
        },
        headers=csrf_headers(client),
    )
    assert regenerate_response.status_code == 400
    assert regenerate_response.json()["detail"] == "Daily menu is closed"

    school_reopen_response = client.post(
        f"/api/v1/menus/weekly/{menu['id']}/days/monday/reopen",
        headers=csrf_headers(client),
    )
    assert school_reopen_response.status_code == 403
    assert school_reopen_response.json()["detail"] == ("Only administrators can reopen daily menus")

    login(client, identities.lower_admin.username, identities.lower_admin_password)

    forbidden_school_response = client.get(
        "/api/v1/menus/weekly/current-week/closed-days",
        params={"school_id": str(identities.other_school.id)},
    )
    assert forbidden_school_response.status_code == 403
    assert forbidden_school_response.json()["detail"] == "School access denied"

    monkeypatch.setattr(
        "app.modules.menus.day_closure._today_in_school_timezone",
        lambda: Date(2026, 7, 13),
    )
    previous_workweek_response = client.post(
        f"/api/v1/menus/weekly/{menu['id']}/days/monday/reopen",
        headers=csrf_headers(client),
    )
    assert previous_workweek_response.status_code == 400
    assert previous_workweek_response.json()["detail"] == (
        "Only current-week daily menus can be reopened"
    )

    monkeypatch.setattr(
        "app.modules.menus.day_closure._today_in_school_timezone",
        lambda: Date(2026, 7, 6),
    )
    closed_days_response = client.get(
        "/api/v1/menus/weekly/current-week/closed-days",
        params={"school_id": str(identities.own_school.id)},
    )
    assert closed_days_response.status_code == 200
    assert closed_days_response.json() == {
        "school_id": str(identities.own_school.id),
        "week_starts_on": "2026-07-06",
        "week_ends_on": "2026-07-10",
        "items": [
            {
                "menu_id": menu["id"],
                "menu_title": menu["title"],
                "meal_type": "lunch",
                "weekday": "monday",
                "date": "2026-07-06",
                "closed_at": closed_day["closed_at"],
                "close_reason": "manual",
            }
        ],
    }

    reopen_response = client.post(
        f"/api/v1/menus/weekly/{menu['id']}/days/monday/reopen",
        headers=csrf_headers(client),
    )
    assert reopen_response.status_code == 200
    reopened_menu = reopen_response.json()
    reopened_day = reopened_menu["days"][0]
    assert reopened_day["closed_at"] is None
    assert reopened_day["closed_by"] is None
    assert reopened_day["close_reason"] is None
    assert reopened_menu["revision"] == closed_menu["revision"] + 1

    duplicate_reopen_response = client.post(
        f"/api/v1/menus/weekly/{menu['id']}/days/monday/reopen",
        headers=csrf_headers(client),
    )
    assert duplicate_reopen_response.status_code == 200
    assert duplicate_reopen_response.json()["revision"] == reopened_menu["revision"]

    closed_days_after_reopen_response = client.get(
        "/api/v1/menus/weekly/current-week/closed-days",
        params={"school_id": str(identities.own_school.id)},
    )
    assert closed_days_after_reopen_response.status_code == 200
    assert closed_days_after_reopen_response.json()["items"] == []

    mongo_client = MongoClient(settings.mongo_uri, tz_aware=True)
    try:
        stored_menu = mongo_client[settings.mongo_db]["weekly_menus"].find_one(
            {"_id": PydanticObjectId(menu["id"])}
        )
    finally:
        mongo_client.close()
    assert stored_menu is not None
    assert stored_menu["days"][0]["close_notification_pending"] is False
    assert stored_menu["days"][0]["reopened_at"] is not None
    assert stored_menu["days"][0]["reopened_by"] == identities.lower_admin.id

    login(client, identities.school_user.username, identities.school_user_password)

    reopened_day["items"][0]["servings"][0]["children_count"] = 6
    reopened_edit_response = client.patch(
        f"/api/v1/menus/weekly/{menu['id']}",
        json={"days": reopened_menu["days"], "revision": reopened_menu["revision"]},
        headers=csrf_headers(client),
    )
    assert reopened_edit_response.status_code == 200
    assert (
        reopened_edit_response.json()["days"][0]["items"][0]["servings"][0]["children_count"] == 6
    )


def test_school_cannot_generate_requirement_for_zero_day(seeded_client) -> None:
    client, identities = seeded_client
    login(client, identities.admin.username, identities.admin_password)

    ingredient_response = client.post(
        "/api/v1/recipes/ingredients",
        json={"name": "Морква", "unit": "g"},
        headers=csrf_headers(client),
    )
    card_id, variant_id = create_confirmed_dish(
        client,
        ingredient_id=ingredient_response.json()["id"],
    )
    menu = publish_school_menu(
        client,
        school_id=str(identities.own_school.id),
        card_id=card_id,
        variant_id=variant_id,
    )

    login(client, identities.school_user.username, identities.school_user_password)
    response = client.post(
        "/api/v1/menu-requirements/generate",
        json={
            "weekly_menu_id": menu["id"],
            "weekday": "monday",
            "service_date": "2026-07-06",
        },
        headers=csrf_headers(client),
    )

    assert response.status_code == 400
    assert response.json()["detail"] == (
        "At least one dish must have a children count greater than zero"
    )


def test_school_added_product_is_isolated_and_included_in_requirement(seeded_client) -> None:
    client, identities = seeded_client
    login(client, identities.admin.username, identities.admin_password)

    carrot_response = client.post(
        "/api/v1/recipes/ingredients",
        json={"name": "Морква", "unit": "g"},
        headers=csrf_headers(client),
    )
    assert carrot_response.status_code == 201
    bread_response = client.post(
        "/api/v1/recipes/ingredients",
        json={"name": "Хліб пшеничний", "unit": "g"},
        headers=csrf_headers(client),
    )
    assert bread_response.status_code == 201
    bread_id = bread_response.json()["id"]
    card_id, variant_id = create_confirmed_dish(
        client,
        ingredient_id=carrot_response.json()["id"],
    )

    create_response = client.post(
        "/api/v1/menus/weekly",
        json={
            "title": "Спільний шаблон",
            "meal_type": "lunch",
            "starts_on": "2026-07-06",
            "ends_on": "2026-07-10",
            "days": [
                {
                    "weekday": "monday",
                    "date": "2026-07-06",
                    "items": [
                        {
                            "position": 1,
                            "kind": "dish_card",
                            "recipe_card_number": "REQ-1",
                            "dish_card_id": card_id,
                            "name": "Овочевий суп",
                            "allergen_codes": [],
                            "portions": [
                                {
                                    "age_group": "6-11",
                                    "yield_amount": "200",
                                    "dish_card_portion_variant_id": variant_id,
                                    "nutrition": {},
                                }
                            ],
                        }
                    ],
                }
            ],
        },
        headers=csrf_headers(client),
    )
    assert create_response.status_code == 201
    template = create_response.json()
    publish_response = client.post(
        f"/api/v1/menus/weekly/{template['id']}/publish",
        json={
            "school_ids": [
                str(identities.own_school.id),
                str(identities.other_school.id),
            ]
        },
        headers=csrf_headers(client),
    )
    assert publish_response.status_code == 200
    own_copy_id, other_copy_id = publish_response.json()["created_menu_ids"]
    own_copy = client.get(f"/api/v1/menus/weekly/{own_copy_id}").json()
    original_item_id = own_copy["days"][0]["items"][0]["id"]
    group = identities.own_school.groups[0]
    own_copy["days"][0]["items"][0]["servings"] = [
        {
            "school_group_id": str(group.id),
            "age_group": group.age_group.value,
            "children_count": 3,
        }
    ]
    own_copy["days"][0]["items"].append(
        {
            "position": 2,
            "kind": "product",
            "product_ingredient_id": bread_id,
            "product_name_snapshot": "Стара назва",
            "name": "Хліб пшеничний",
            "allergen_codes": [],
            "portions": [
                {
                    "age_group": group.age_group.value,
                    "yield_amount": "30",
                    "nutrition": {},
                }
            ],
            "servings": [
                {
                    "school_group_id": str(group.id),
                    "age_group": group.age_group.value,
                    "children_count": 3,
                }
            ],
        }
    )

    login(client, identities.school_user.username, identities.school_user_password)
    update_response = client.patch(
        f"/api/v1/menus/weekly/{own_copy_id}",
        json={"days": own_copy["days"], "revision": own_copy["revision"]},
        headers=csrf_headers(client),
    )
    assert update_response.status_code == 200, update_response.text
    updated_menu = update_response.json()
    assert len(updated_menu["days"][0]["items"]) == 2
    assert updated_menu["days"][0]["items"][0]["id"] == original_item_id
    added_item = updated_menu["days"][0]["items"][1]
    assert added_item["id"] != original_item_id
    assert added_item["product_ingredient_id"] == bread_id
    assert added_item["product_name_snapshot"] == "Хліб пшеничний"

    own_get_response = client.get(f"/api/v1/menus/weekly/{own_copy_id}")
    assert own_get_response.status_code == 200
    assert [item["name"] for item in own_get_response.json()["days"][0]["items"]] == [
        "Овочевий суп",
        "Хліб пшеничний",
    ]

    generate_response = client.post(
        "/api/v1/menu-requirements/generate",
        json={
            "weekly_menu_id": own_copy_id,
            "weekday": "monday",
            "service_date": "2026-07-06",
        },
        headers=csrf_headers(client),
    )
    assert generate_response.status_code == 200, generate_response.text
    requirement = generate_response.json()["items"][0]
    bread_dish = next(dish for dish in requirement["dishes"] if dish["name"] == "Хліб пшеничний")
    assert bread_dish["menu_item_id"] == added_item["id"]
    bread_row = next(
        row for row in requirement["ingredient_rows"] if row["ingredient_id"] == bread_id
    )
    assert bread_row["per_person_total_g"] == "30"
    assert bread_row["issue_total_raw_g"] == "90"

    login(client, identities.admin.username, identities.admin_password)
    template_after = client.get(f"/api/v1/menus/weekly/{template['id']}").json()
    other_copy_after = client.get(f"/api/v1/menus/weekly/{other_copy_id}").json()
    assert [item["name"] for item in template_after["days"][0]["items"]] == ["Овочевий суп"]
    assert [item["name"] for item in other_copy_after["days"][0]["items"]] == [
        "Овочевий суп"
    ]

    requests_response = client.get("/api/v1/menus/change-requests")
    assert requests_response.status_code == 200
    added_change = next(
        change
        for change in requests_response.json()["items"][0]["changes"]
        if change["field"] == "item_added"
    )
    assert added_change["item_id"] == added_item["id"]
    assert added_change["after_value"] == "Хліб пшеничний"


def test_generate_menu_requirement_for_wholegrain_bread_product_composite_yield(
    seeded_client,
) -> None:
    client, identities = seeded_client
    login(client, identities.admin.username, identities.admin_password)

    bread_resp = client.post(
        "/api/v1/recipes/ingredients",
        json={"name": "Хліб цільнозерновий", "unit": "g"},
        headers=csrf_headers(client),
    )
    assert bread_resp.status_code == 201
    bread_id = bread_resp.json()["id"]

    create_response = client.post(
        "/api/v1/menus/weekly",
        json={
            "title": "Меню з цільнозерновим хлібом",
            "meal_type": "lunch",
            "starts_on": "2026-07-06",
            "ends_on": "2026-07-10",
            "days": [
                {
                    "weekday": "monday",
                    "date": "2026-07-06",
                    "items": [
                        {
                            "position": 1,
                            "kind": "product",
                            "product_ingredient_id": bread_id,
                            "product_name_snapshot": "Хліб цільнозерновий",
                            "name": "Хліб цільнозерновий",
                            "allergen_codes": [],
                            "portions": [
                                {
                                    "age_group": "6-11",
                                    "yield_amount": "20/20",
                                    "nutrition": {},
                                }
                            ],
                        }
                    ],
                }
            ],
        },
        headers=csrf_headers(client),
    )
    assert create_response.status_code == 201
    template = create_response.json()

    publish_response = client.post(
        f"/api/v1/menus/weekly/{template['id']}/publish",
        json={"school_ids": [str(identities.own_school.id)]},
        headers=csrf_headers(client),
    )
    assert publish_response.status_code == 200
    own_copy_id = publish_response.json()["created_menu_ids"][0]
    own_copy = client.get(f"/api/v1/menus/weekly/{own_copy_id}").json()

    group = identities.own_school.groups[0]
    children_count = 12
    own_copy["days"][0]["items"][0]["servings"] = [
        {
            "school_group_id": str(group.id),
            "age_group": group.age_group.value,
            "children_count": children_count,
        }
    ]

    login(client, identities.school_user.username, identities.school_user_password)
    update_response = client.patch(
        f"/api/v1/menus/weekly/{own_copy_id}",
        json={"days": own_copy["days"], "revision": own_copy["revision"]},
        headers=csrf_headers(client),
    )
    assert update_response.status_code == 200, update_response.text

    generate_response = client.post(
        "/api/v1/menu-requirements/generate",
        json={
            "weekly_menu_id": own_copy_id,
            "weekday": "monday",
            "service_date": "2026-07-06",
        },
        headers=csrf_headers(client),
    )
    assert generate_response.status_code == 200, generate_response.text
    requirement = generate_response.json()["items"][0]

    bread_dish = next(
        dish for dish in requirement["dishes"] if dish["name"] == "Хліб цільнозерновий"
    )
    assert bread_dish["children_count"] == children_count
    assert bread_dish["yield_amount"] == "20/20"

    bread_row = next(
        row for row in requirement["ingredient_rows"] if row["ingredient_id"] == bread_id
    )
    assert bread_row["ingredient_name"] == "Хліб цільнозерновий"
    assert bread_row["per_person_total_g"] == "40"
    assert bread_row["issue_total_raw_g"] == "480"
    assert bread_row["gross_per_person_total_g"] == "40"
    assert bread_row["gross_issue_total_rounded_g"] == 480

    cell = bread_row["cells"][0]
    assert cell["net_per_person_g"] == "40"
    assert cell["gross_per_person_g"] == "40"

    updated_menu = update_response.json()
    updated_menu["days"][0]["items"][0]["portions"][0]["yield_amount"] = "40 г"
    patch2_resp = client.patch(
        f"/api/v1/menus/weekly/{own_copy_id}",
        json={"days": updated_menu["days"], "revision": updated_menu["revision"]},
        headers=csrf_headers(client),
    )
    assert patch2_resp.status_code == 200

    regen_resp = client.post(
        "/api/v1/menu-requirements/generate",
        json={
            "weekly_menu_id": own_copy_id,
            "weekday": "monday",
            "service_date": "2026-07-06",
        },
        headers=csrf_headers(client),
    )
    assert regen_resp.status_code == 200
    regen_req = regen_resp.json()["items"][0]
    regen_row = next(
        row for row in regen_req["ingredient_rows"] if row["ingredient_id"] == bread_id
    )
    assert regen_row["per_person_total_g"] == "40"
    assert regen_row["issue_total_raw_g"] == "480"

    invalid_menu = patch2_resp.json()
    invalid_menu["days"][0]["items"][0]["portions"][0]["yield_amount"] = "невідомо"
    patch3_resp = client.patch(
        f"/api/v1/menus/weekly/{own_copy_id}",
        json={"days": invalid_menu["days"], "revision": invalid_menu["revision"]},
        headers=csrf_headers(client),
    )
    assert patch3_resp.status_code == 200

    invalid_gen_resp = client.post(
        "/api/v1/menu-requirements/generate",
        json={
            "weekly_menu_id": own_copy_id,
            "weekday": "monday",
            "service_date": "2026-07-06",
        },
        headers=csrf_headers(client),
    )
    assert invalid_gen_resp.status_code == 400
    assert invalid_gen_resp.json()["detail"] == (
        'Product "Хліб цільнозерновий" must have a single numeric yield in grams'
    )

