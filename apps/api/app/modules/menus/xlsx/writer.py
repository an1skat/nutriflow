import re
from datetime import UTC, datetime
from decimal import Decimal

from app.modules.menus.models import MenuItemKind, Weekday, WeeklyMenu
from app.modules.menus.xlsx.common import (
    AGE_GROUP_HEADERS,
    AGE_GROUPS_BY_BLOCK,
    BASE_HEADERS,
    DAY_LABELS,
    DAY_TOTAL_LABEL,
    FIRST_CONTENT_ROW,
    HEADER_GROUP_ROW,
    HEADER_SUBHEADER_ROW,
    META_SHEET_NAME,
    OFFICIAL_PORTION_BLOCK_STARTS,
    PORTION_SUBHEADERS,
    SCHEMA_VERSION,
    TEMPLATE_SHEET_TITLES,
    VISIBLE_MAX_COLUMN,
    WEEK_TITLE_ROW,
    XlsxMenuError,
    import_openpyxl,
    save_workbook,
    school_weekdays,
)

TEMPLATE_EMPTY_ITEM_ROWS = 6


def build_template_workbook(*, weeks: int = 4) -> bytes:
    if weeks < 1 or weeks > 4:
        raise XlsxMenuError("Template can contain from one to four weeks")

    openpyxl = import_openpyxl()
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    for cycle_week, sheet_name in enumerate(TEMPLATE_SHEET_TITLES[:weeks], start=1):
        sheet = workbook.create_sheet(sheet_name)
        _write_sheet_headers(sheet, cycle_week=cycle_week)
        row_number = FIRST_CONTENT_ROW
        for weekday in school_weekdays():
            _write_day_separator(sheet, row_number, weekday)
            first_item_row = row_number + 1
            total_row = first_item_row + TEMPLATE_EMPTY_ITEM_ROWS
            _write_total_row(sheet, total_row, first_item_row, total_row - 1)
            row_number = total_row + 1
        _apply_print_settings(sheet, row_number - 1)

    _write_meta_sheet(workbook, [])
    return save_workbook(workbook)


def build_export_workbook(menus: list[WeeklyMenu]) -> bytes:
    if not menus:
        raise XlsxMenuError("At least one weekly menu is required for export")
    if len(menus) > 4:
        raise XlsxMenuError("Can export at most four weekly menus in one workbook")

    openpyxl = import_openpyxl()
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    used_titles: set[str] = set()
    for index, menu in enumerate(menus, start=1):
        cycle_week = menu.cycle_week or index
        sheet_title = _unique_sheet_title(_week_sheet_title(cycle_week), used_titles)
        used_titles.add(sheet_title)
        sheet = workbook.create_sheet(sheet_title)
        _write_sheet_headers(sheet, cycle_week=cycle_week)
        _write_menu_rows(sheet, menu)

    _write_meta_sheet(workbook, menus)
    return save_workbook(workbook)


def _write_sheet_headers(sheet, *, cycle_week: int) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    green = "166534"
    pale_green = "DCFCE7"
    light_green = "F0FDF4"
    border_side = Side(style="thin", color="94A3B8")
    border = Border(
        left=border_side,
        right=border_side,
        top=border_side,
        bottom=border_side,
    )

    for cell_range in ("A1:A2", "B1:B2", "C1:C2", "D1:H1", "I1:M1", "N1:R1"):
        sheet.merge_cells(cell_range)

    for column_number, header in enumerate(BASE_HEADERS, start=1):
        sheet.cell(HEADER_GROUP_ROW, column_number, header)
    for start_col, header in zip(
        OFFICIAL_PORTION_BLOCK_STARTS,
        AGE_GROUP_HEADERS,
        strict=True,
    ):
        sheet.cell(HEADER_GROUP_ROW, start_col, header)
        for offset, subheader in enumerate(PORTION_SUBHEADERS):
            sheet.cell(HEADER_SUBHEADER_ROW, start_col + offset, subheader)

    for row in sheet.iter_rows(
        min_row=HEADER_GROUP_ROW,
        max_row=HEADER_SUBHEADER_ROW,
        min_col=1,
        max_col=VISIBLE_MAX_COLUMN,
    ):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.font = Font(bold=True, color="FFFFFF" if cell.row == 1 else "14532D")
            cell.fill = PatternFill(
                "solid",
                fgColor=green if cell.row == 1 else light_green,
            )

    sheet.merge_cells(
        start_row=WEEK_TITLE_ROW,
        start_column=1,
        end_row=WEEK_TITLE_ROW,
        end_column=VISIBLE_MAX_COLUMN,
    )
    week_cell = sheet.cell(WEEK_TITLE_ROW, 1, f"{cycle_week}-й тиждень")
    week_cell.font = Font(bold=True, size=13, color="14532D")
    week_cell.fill = PatternFill("solid", fgColor=pale_green)
    week_cell.alignment = Alignment(horizontal="center", vertical="center")
    week_cell.border = border

    widths = [30, 16, 42, 12, 16, 11, 11, 14, 12, 16, 11, 11, 14, 12, 16, 11, 11, 14]
    for column_number, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column_number)].width = width

    sheet.row_dimensions[1].height = 42
    sheet.row_dimensions[2].height = 50
    sheet.row_dimensions[3].height = 24
    sheet.freeze_panes = "D5"


def _write_day_separator(sheet, row_number: int, weekday: Weekday) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    side = Side(style="thin", color="94A3B8")
    for column_number in range(1, VISIBLE_MAX_COLUMN + 1):
        cell = sheet.cell(row_number, column_number)
        cell.fill = PatternFill("solid", fgColor="DCFCE7")
        cell.border = Border(left=side, right=side, top=side, bottom=side)
    cell = sheet.cell(row_number, 3, DAY_LABELS[weekday])
    cell.font = Font(bold=True, color="14532D")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[row_number].height = 24


def _write_total_row(
    sheet,
    row_number: int,
    first_item_row: int,
    last_item_row: int,
) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    side = Side(style="thin", color="94A3B8")
    for column_number in range(1, VISIBLE_MAX_COLUMN + 1):
        cell = sheet.cell(row_number, column_number)
        cell.fill = PatternFill("solid", fgColor="FEF3C7")
        cell.border = Border(left=side, right=side, top=side, bottom=side)
        cell.font = Font(bold=True, color="713F12")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.cell(row_number, 3, DAY_TOTAL_LABEL)
    for column_number in range(4, VISIBLE_MAX_COLUMN + 1):
        column_letter = get_column_letter(column_number)
        sheet.cell(
            row_number,
            column_number,
            f"=SUM({column_letter}{first_item_row}:{column_letter}{last_item_row})",
        )
        sheet.cell(row_number, column_number).number_format = "0.00"


def _write_menu_rows(sheet, menu: WeeklyMenu) -> None:
    from openpyxl.styles import Alignment, Border, Side

    row_number = FIRST_CONTENT_ROW
    side = Side(style="thin", color="CBD5E1")
    border = Border(left=side, right=side, top=side, bottom=side)

    for day in sorted(menu.days, key=lambda value: list(Weekday).index(value.weekday)):
        _write_day_separator(sheet, row_number, day.weekday)
        first_item_row = row_number + 1
        row_number = first_item_row

        for item in sorted(day.items, key=lambda value: value.position):
            sheet.cell(row_number, 1, _export_source_text(item))
            sheet.cell(row_number, 2, ", ".join(item.allergen_codes))
            sheet.cell(row_number, 3, item.name)

            portions_by_age = {portion.age_group: portion for portion in item.portions}
            for age_group, start_col in zip(
                AGE_GROUPS_BY_BLOCK,
                OFFICIAL_PORTION_BLOCK_STARTS,
                strict=True,
            ):
                portion = portions_by_age.get(age_group)
                if portion is None:
                    continue
                sheet.cell(row_number, start_col, portion.yield_amount)
                sheet.cell(row_number, start_col + 1, _decimal_for_excel(portion.nutrition.kcal))
                sheet.cell(
                    row_number,
                    start_col + 2,
                    _decimal_for_excel(portion.nutrition.proteins),
                )
                sheet.cell(row_number, start_col + 3, _decimal_for_excel(portion.nutrition.fats))
                sheet.cell(row_number, start_col + 4, _decimal_for_excel(portion.nutrition.carbs))

            for column_number in range(1, VISIBLE_MAX_COLUMN + 1):
                cell = sheet.cell(row_number, column_number)
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            for column_number in range(5, VISIBLE_MAX_COLUMN + 1):
                if column_number not in OFFICIAL_PORTION_BLOCK_STARTS:
                    sheet.cell(row_number, column_number).number_format = "0.00"
            sheet.row_dimensions[row_number].height = 34
            row_number += 1

        _write_total_row(sheet, row_number, first_item_row, row_number - 1)
        row_number += 1

    _apply_print_settings(sheet, row_number - 1)


def _write_meta_sheet(workbook, menus: list[WeeklyMenu]) -> None:
    meta = workbook.create_sheet(META_SHEET_NAME)
    meta.append(["key", "value"])
    meta.append(["schema_version", SCHEMA_VERSION])
    meta.append(["document_type", "weekly_menu_batch"])
    meta.append(["generated_at", datetime.now(UTC).isoformat()])
    meta.append([])
    meta.append(
        [
            "sheet_name",
            "cycle_week",
            "menu_id",
            "school_id",
            "source_menu_id",
            "meal_type",
        ]
    )
    for index, menu in enumerate(menus, start=1):
        meta.append(
            [
                workbook.worksheets[index - 1].title,
                menu.cycle_week,
                str(menu.id or ""),
                str(menu.school_id or ""),
                str(menu.source_menu_id or ""),
                menu.meal_type.value,
            ]
        )
    meta.sheet_state = "hidden"


def _apply_print_settings(sheet, last_row: int) -> None:
    sheet.auto_filter.ref = f"A1:R{max(last_row, 3)}"
    sheet.print_area = f"A1:R{max(last_row, 3)}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"
    sheet.page_margins.left = 0.2
    sheet.page_margins.right = 0.2
    sheet.page_margins.top = 0.3
    sheet.page_margins.bottom = 0.3


def _export_source_text(item) -> str:
    if item.source_text:
        return item.source_text
    if item.kind == MenuItemKind.PRODUCT:
        return "пром. вироб."
    if item.recipe_card_number:
        return f"ТК № {item.recipe_card_number}"
    return ""


def _decimal_for_excel(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


def _week_sheet_title(cycle_week: int) -> str:
    if 1 <= cycle_week <= len(TEMPLATE_SHEET_TITLES):
        return TEMPLATE_SHEET_TITLES[cycle_week - 1]
    return f"Тиждень {cycle_week}"


def _unique_sheet_title(preferred: str, used_titles: set[str]) -> str:
    sanitized = re.sub(r'[:\\/?*\[\]]', "-", preferred).strip() or "Меню"
    sanitized = sanitized[:31]
    if sanitized not in used_titles:
        return sanitized
    suffix = 2
    while True:
        candidate = f"{sanitized[: 28 - len(str(suffix))]} ({suffix})"
        if candidate not in used_titles:
            return candidate
        suffix += 1
