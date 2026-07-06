import re
from copy import deepcopy
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from beanie import PydanticObjectId
from fastapi import UploadFile

from app.modules.auth.service import user_has_permissions
from app.modules.identity.models import AdminPermission, AgeGroup, School, User, UserRole
from app.modules.menus.models import (
    DailyMenu,
    DailyMenuItem,
    MealType,
    MenuItemKind,
    MenuItemServingCount,
    MenuNutrition,
    MenuPortion,
    Weekday,
    WeeklyMenu,
    WeeklyMenuStatus,
)
from app.modules.menus.schemas import (
    CreateWeeklyMenuRequest,
    DailyMenuItemPayload,
    DailyMenuPayload,
    MenuItemServingCountPayload,
    MenuPortionPayload,
    PublishWeeklyMenuRequest,
    PublishWeeklyMenuResponse,
    UpdateWeeklyMenuRequest,
    WeeklyMenuImportPreviewResponse,
)
from app.modules.recipe.models import DishCard, DishCardVersion, Ingredient


class MenuNotFoundError(ValueError):
    """The requested menu does not exist."""


class MenuAccessDeniedError(ValueError):
    """Current user cannot access the requested menu."""


class MenuValidationError(ValueError):
    """Menu payload violates a domain rule."""


class MenuImportError(ValueError):
    """Menu import file cannot be parsed into a weekly menu."""


UKRAINIAN_WEEKDAYS = {
    "понеділок": Weekday.MONDAY,
    "вівторок": Weekday.TUESDAY,
    "середа": Weekday.WEDNESDAY,
    "четвер": Weekday.THURSDAY,
    "п'ятниця": Weekday.FRIDAY,
    "п’ятниця": Weekday.FRIDAY,
    "пятниця": Weekday.FRIDAY,
    "субота": Weekday.SATURDAY,
    "неділя": Weekday.SUNDAY,
}

AGE_GROUPS_BY_BLOCK = [
    AgeGroup.SIX_TO_ELEVEN,
    AgeGroup.ELEVEN_TO_FOURTEEN,
    AgeGroup.FOURTEEN_TO_EIGHTEEN,
]


async def list_weekly_menus(
    current_user: User,
    *,
    offset: int,
    limit: int,
    school_id: PydanticObjectId | None = None,
    template_only: bool = False,
    status: WeeklyMenuStatus | None = None,
    meal_type: MealType | None = None,
) -> tuple[list[WeeklyMenu], int]:
    filters: dict[str, Any] = {}

    if current_user.role == UserRole.SCHOOL_USER:
        filters["school_id"] = current_user.school_id
    elif current_user.role == UserRole.ADMIN:
        await _ensure_menu_permission(current_user)

        if template_only:
            filters["school_id"] = None
            filters["created_by"] = current_user.id
        elif school_id is not None:
            await _ensure_admin_school_access(current_user, school_id)
            filters["school_id"] = school_id
        else:
            owned_school_ids = await _get_admin_school_ids(current_user)
            filters["$or"] = [
                {"school_id": {"$in": owned_school_ids}},
                {"school_id": None, "created_by": current_user.id},
            ]
    elif template_only:
        filters["school_id"] = None
    elif school_id is not None:
        filters["school_id"] = school_id

    if status is not None:
        filters["status"] = status.value
    if meal_type is not None:
        filters["meal_type"] = meal_type.value

    cursor = WeeklyMenu.find(filters)
    total = await cursor.count()
    items = await cursor.sort("-created_at").skip(offset).limit(limit).to_list()
    return items, total


async def get_weekly_menu(
    menu_id: PydanticObjectId,
    current_user: User,
) -> WeeklyMenu:
    menu = await WeeklyMenu.get(menu_id)

    if menu is None:
        raise MenuNotFoundError("Weekly menu not found")

    await _authorize_menu_access(menu, current_user)
    return menu


async def create_weekly_menu(
    data: CreateWeeklyMenuRequest,
    *,
    current_user: User,
) -> WeeklyMenu:
    await _ensure_menu_permission(current_user)

    if data.school_id is not None:
        await _get_active_school_for_user(data.school_id, current_user)

    now = datetime.now(UTC)
    menu = WeeklyMenu(
        title=data.title,
        school_id=data.school_id,
        meal_type=data.meal_type,
        cycle_week=data.cycle_week,
        starts_on=data.starts_on,
        ends_on=data.ends_on,
        days=await _to_daily_menus(data.days),
        notes=data.notes,
        source_file_name=data.source_file_name,
        source_sheet_name=data.source_sheet_name,
        created_by=current_user.id,
        updated_by=current_user.id,
        created_at=now,
        updated_at=now,
    )
    await menu.insert()
    return menu


async def update_weekly_menu(
    menu_id: PydanticObjectId,
    data: UpdateWeeklyMenuRequest,
    current_user: User,
) -> WeeklyMenu:
    menu = await get_weekly_menu(menu_id, current_user)

    if current_user.role == UserRole.SCHOOL_USER and "days" in data.model_fields_set:
        _ensure_school_menu_shape_is_stable(menu, data.days or [])

    if "title" in data.model_fields_set:
        menu.title = data.title
    if "meal_type" in data.model_fields_set:
        menu.meal_type = data.meal_type
    if "cycle_week" in data.model_fields_set:
        menu.cycle_week = data.cycle_week
    if "starts_on" in data.model_fields_set:
        menu.starts_on = data.starts_on
    if "ends_on" in data.model_fields_set:
        menu.ends_on = data.ends_on
    if "days" in data.model_fields_set:
        menu.days = await _to_daily_menus(data.days or [])
    if "notes" in data.model_fields_set:
        menu.notes = data.notes

    menu.updated_by = current_user.id
    menu.updated_at = datetime.now(UTC)
    await menu.save()
    return menu


async def publish_weekly_menu(
    menu_id: PydanticObjectId,
    data: PublishWeeklyMenuRequest,
    admin: User,
) -> PublishWeeklyMenuResponse:
    source = await get_weekly_menu(menu_id, admin)

    if source.school_id is not None:
        raise MenuValidationError("Only template weekly menus can be published")

    target_schools = await _get_publish_target_schools(data.school_ids, admin)
    created_menu_ids: list[PydanticObjectId] = []
    replaced_menu_ids: list[PydanticObjectId] = []
    skipped_existing_school_ids: list[PydanticObjectId] = []
    now = datetime.now(UTC)

    source.status = WeeklyMenuStatus.PUBLISHED
    source.published_at = now
    source.updated_by = admin.id
    source.updated_at = now
    await source.save()

    for school in target_schools:
        existing = await WeeklyMenu.find_one(
            WeeklyMenu.source_menu_id == source.id,
            WeeklyMenu.school_id == school.id,
        )

        if existing is not None and not data.replace_existing:
            skipped_existing_school_ids.append(school.id)
            continue

        if existing is not None:
            existing.title = source.title
            existing.meal_type = source.meal_type
            existing.cycle_week = source.cycle_week
            existing.starts_on = source.starts_on
            existing.ends_on = source.ends_on
            existing.status = WeeklyMenuStatus.PUBLISHED
            existing.days = deepcopy(source.days)
            existing.notes = source.notes
            existing.source_file_name = source.source_file_name
            existing.source_sheet_name = source.source_sheet_name
            existing.published_at = now
            existing.updated_by = admin.id
            existing.updated_at = now
            await existing.save()
            replaced_menu_ids.append(existing.id)
            continue

        copy = WeeklyMenu(
            title=source.title,
            school_id=school.id,
            source_menu_id=source.id,
            meal_type=source.meal_type,
            cycle_week=source.cycle_week,
            starts_on=source.starts_on,
            ends_on=source.ends_on,
            status=WeeklyMenuStatus.PUBLISHED,
            days=deepcopy(source.days),
            notes=source.notes,
            source_file_name=source.source_file_name,
            source_sheet_name=source.source_sheet_name,
            published_at=now,
            created_by=admin.id,
            updated_by=admin.id,
            created_at=now,
            updated_at=now,
        )
        await copy.insert()
        created_menu_ids.append(copy.id)

    return PublishWeeklyMenuResponse(
        source_menu_id=source.id,
        target_school_ids=[school.id for school in target_schools],
        created_menu_ids=created_menu_ids,
        replaced_menu_ids=replaced_menu_ids,
        skipped_existing_school_ids=skipped_existing_school_ids,
    )


async def preview_weekly_menu_import(
    file: UploadFile,
    *,
    meal_type: MealType,
    sheet_name: str | None,
    title: str | None,
) -> WeeklyMenuImportPreviewResponse:
    content = await file.read()
    menu, warnings, parsed_sheet_name = parse_weekly_menu_workbook(
        content,
        filename=file.filename or "",
        meal_type=meal_type,
        sheet_name=sheet_name,
        title=title,
    )
    return WeeklyMenuImportPreviewResponse(
        filename=file.filename or "",
        sheet_name=parsed_sheet_name,
        warnings=warnings,
        menu=menu,
    )


async def create_weekly_menu_from_import(
    file: UploadFile,
    *,
    meal_type: MealType,
    sheet_name: str | None,
    title: str | None,
    school_id: PydanticObjectId | None,
    current_user: User,
) -> WeeklyMenu:
    content = await file.read()
    menu_request, _warnings, _parsed_sheet_name = parse_weekly_menu_workbook(
        content,
        filename=file.filename or "",
        meal_type=meal_type,
        sheet_name=sheet_name,
        title=title,
    )
    menu_request.school_id = school_id
    return await create_weekly_menu(menu_request, current_user=current_user)


def parse_weekly_menu_workbook(
    content: bytes,
    *,
    filename: str,
    meal_type: MealType,
    sheet_name: str | None = None,
    title: str | None = None,
) -> tuple[CreateWeeklyMenuRequest, list[str], str]:
    try:
        import openpyxl
    except ModuleNotFoundError as exc:
        raise MenuImportError("openpyxl is required to parse menu workbooks") from exc

    try:
        workbook = openpyxl.load_workbook(
            BytesIO(content),
            data_only=True,
            read_only=True,
        )
    except Exception as exc:
        raise MenuImportError(f"Could not read .xlsx file: {exc}") from exc

    if not workbook.sheetnames:
        raise MenuImportError("Workbook does not contain sheets")

    selected_sheet_name = sheet_name or workbook.sheetnames[0]
    if selected_sheet_name not in workbook.sheetnames:
        raise MenuImportError("Requested sheet was not found in workbook")

    sheet = workbook[selected_sheet_name]
    source_col, allergen_col, name_col = _detect_menu_columns(sheet)
    portion_blocks = _detect_portion_blocks(sheet, name_col)
    day_rows = _detect_day_rows(sheet, name_col)

    if len(portion_blocks) < 3:
        raise MenuImportError("Could not detect all age-group nutrition blocks")
    if not day_rows:
        raise MenuImportError("Could not detect weekday blocks")

    warnings: list[str] = []
    days: list[DailyMenuPayload] = []
    week_number = _detect_cycle_week(sheet, selected_sheet_name)

    for index, (row_number, weekday) in enumerate(day_rows):
        next_row = day_rows[index + 1][0] if index + 1 < len(day_rows) else sheet.max_row + 1
        items = _parse_day_items(
            sheet,
            start_row=row_number + 1,
            end_row=next_row,
            source_col=source_col,
            allergen_col=allergen_col,
            name_col=name_col,
            portion_blocks=portion_blocks[:3],
            warnings=warnings,
        )
        if items:
            days.append(DailyMenuPayload(weekday=weekday, items=items))

    if not days:
        raise MenuImportError("Workbook sheet does not contain menu items")

    resolved_title = title or _default_import_title(
        filename=filename,
        sheet_name=selected_sheet_name,
        meal_type=meal_type,
    )

    return (
        CreateWeeklyMenuRequest(
            title=resolved_title,
            meal_type=meal_type,
            cycle_week=week_number,
            days=days,
            source_file_name=filename,
            source_sheet_name=selected_sheet_name,
        ),
        warnings,
        selected_sheet_name,
    )


async def _to_daily_menus(data: list[DailyMenuPayload]) -> list[DailyMenu]:
    days: list[DailyMenu] = []

    for day in data:
        items = [await _to_daily_menu_item(item) for item in day.items]
        days.append(
            DailyMenu(
                weekday=day.weekday,
                date=day.date,
                items=sorted(items, key=lambda item: item.position),
                notes=day.notes,
            )
        )

    return sorted(days, key=lambda day: list(Weekday).index(day.weekday))


async def _to_daily_menu_item(data: DailyMenuItemPayload) -> DailyMenuItem:
    item = DailyMenuItem(
        id=data.id or PydanticObjectId(),
        position=data.position,
        kind=data.kind,
        source_text=data.source_text,
        recipe_card_number=data.recipe_card_number,
        dish_card_id=data.dish_card_id,
        dish_card_version_id=data.dish_card_version_id,
        product_ingredient_id=data.product_ingredient_id,
        product_name_snapshot=data.product_name_snapshot,
        name=data.name,
        allergen_codes=data.allergen_codes,
        portions=[_to_menu_portion(portion) for portion in data.portions],
        servings=[_to_serving_count(serving) for serving in data.servings],
        notes=data.notes,
    )
    await _resolve_item_references(item)
    return item


def _to_menu_portion(data: MenuPortionPayload) -> MenuPortion:
    return MenuPortion(
        age_group=data.age_group,
        yield_amount=data.yield_amount,
        dish_card_portion_variant_id=data.dish_card_portion_variant_id,
        nutrition=MenuNutrition(**data.nutrition.model_dump()),
    )


def _to_serving_count(data: MenuItemServingCountPayload) -> MenuItemServingCount:
    return MenuItemServingCount(
        school_group_id=data.school_group_id,
        age_group=data.age_group,
        children_count=data.children_count,
    )


async def _resolve_item_references(item: DailyMenuItem) -> None:
    if item.kind == MenuItemKind.PRODUCT:
        if item.product_ingredient_id is not None:
            ingredient = await Ingredient.get(item.product_ingredient_id)
            if ingredient is None:
                raise MenuValidationError("Product ingredient not found")
            item.product_name_snapshot = item.product_name_snapshot or ingredient.name
        return

    if item.dish_card_id is None and item.recipe_card_number:
        dish_card = await _find_dish_card_by_number(item.recipe_card_number)
        if dish_card is not None:
            item.dish_card_id = dish_card.id
            item.dish_card_version_id = item.dish_card_version_id or dish_card.current_version_id

    if item.dish_card_id is None:
        return

    dish_card = await DishCard.get(item.dish_card_id)
    if dish_card is None:
        raise MenuValidationError("Dish card not found")

    if item.dish_card_version_id is None:
        item.dish_card_version_id = dish_card.current_version_id
        return

    version = await DishCardVersion.get(item.dish_card_version_id)
    if version is None or version.dish_card_id != dish_card.id:
        raise MenuValidationError("Dish card version does not belong to menu item dish card")


async def _find_dish_card_by_number(card_number: str) -> DishCard | None:
    for candidate in _card_number_candidates(card_number):
        dish_card = await DishCard.find_one(DishCard.card_number == candidate)
        if dish_card is not None:
            return dish_card
    return None


def _card_number_candidates(card_number: str) -> list[str]:
    normalized = card_number.strip()
    candidates = [normalized]
    without_leading_zero = re.sub(r"\b0+(\d)", r"\1", normalized)
    if without_leading_zero not in candidates:
        candidates.append(without_leading_zero)
    return candidates


async def _authorize_menu_access(menu: WeeklyMenu, current_user: User) -> None:
    if current_user.role == UserRole.OWNER:
        return

    if current_user.role == UserRole.ADMIN:
        await _ensure_menu_permission(current_user)

        if menu.school_id is None:
            if menu.created_by == current_user.id:
                return
            raise MenuAccessDeniedError("Menu access denied")

        await _ensure_admin_school_access(current_user, menu.school_id)
        return

    if menu.school_id != current_user.school_id:
        raise MenuAccessDeniedError("School access denied")


async def _get_active_school(school_id: PydanticObjectId) -> School:
    school = await School.get(school_id)
    if school is None:
        raise MenuValidationError("School not found")
    if not school.is_active:
        raise MenuValidationError("School is inactive")
    return school


async def _get_active_school_for_user(
    school_id: PydanticObjectId,
    current_user: User,
) -> School:
    school = await _get_active_school(school_id)

    if current_user.role == UserRole.ADMIN and school.admin_owner_id != current_user.id:
        raise MenuAccessDeniedError("School access denied")

    return school


async def _get_publish_target_schools(
    school_ids: list[PydanticObjectId] | None,
    current_user: User,
) -> list[School]:
    if school_ids is None:
        if current_user.role == UserRole.ADMIN:
            return (
                await School.find(
                    School.is_active == True,  # noqa: E712
                    School.admin_owner_id == current_user.id,
                )
                .sort("name")
                .to_list()
            )

        return await School.find(School.is_active == True).sort("name").to_list()  # noqa: E712

    schools: list[School] = []
    seen: set[PydanticObjectId] = set()
    for school_id in school_ids:
        if school_id in seen:
            continue
        seen.add(school_id)
        schools.append(await _get_active_school_for_user(school_id, current_user))
    return schools


async def _ensure_menu_permission(current_user: User) -> None:
    if not await user_has_permissions(current_user, AdminPermission.MENUS_MANAGE):
        raise MenuAccessDeniedError("Insufficient permissions")


async def _ensure_admin_school_access(
    current_user: User,
    school_id: PydanticObjectId,
) -> None:
    school = await School.get(school_id)

    if school is None:
        raise MenuValidationError("School not found")
    if current_user.role == UserRole.ADMIN and school.admin_owner_id != current_user.id:
        raise MenuAccessDeniedError("School access denied")


async def _get_admin_school_ids(current_user: User) -> list[PydanticObjectId]:
    schools = await School.find(School.admin_owner_id == current_user.id).to_list()
    return [school.id for school in schools if school.id is not None]


def _ensure_school_menu_shape_is_stable(
    current_menu: WeeklyMenu,
    new_days: list[DailyMenuPayload],
) -> None:
    current_counts = {day.weekday: len(day.items) for day in current_menu.days}
    new_counts = {day.weekday: len(day.items) for day in new_days}

    if current_counts != new_counts:
        raise MenuValidationError("School users cannot change the number of dishes in a day")


def _detect_menu_columns(sheet) -> tuple[int, int, int]:
    source_col = allergen_col = name_col = None

    for column, cell in enumerate(sheet[1], start=1):
        value = _cell_text(cell.value).lower()
        if "збірник" in value or "розкладки" in value:
            source_col = column
        elif "алерген" in value:
            allergen_col = column
        elif "найменування" in value:
            name_col = column

    if source_col is None or allergen_col is None or name_col is None:
        raise MenuImportError("Could not detect source, allergen, and dish name columns")

    return source_col, allergen_col, name_col


def _detect_portion_blocks(sheet, name_col: int) -> list[int]:
    blocks: list[int] = []
    for column, cell in enumerate(sheet[2], start=1):
        if column <= name_col:
            continue
        value = _cell_text(cell.value).lower()
        if value.startswith("вихід"):
            blocks.append(column)
    return blocks


def _detect_day_rows(sheet, name_col: int) -> list[tuple[int, Weekday]]:
    day_rows: list[tuple[int, Weekday]] = []

    for row_number in range(1, sheet.max_row + 1):
        for column in range(max(1, name_col - 1), name_col + 2):
            value = _cell_text(sheet.cell(row_number, column).value).lower()
            if value in UKRAINIAN_WEEKDAYS:
                day_rows.append((row_number, UKRAINIAN_WEEKDAYS[value]))
                break

    return day_rows


def _parse_day_items(
    sheet,
    *,
    start_row: int,
    end_row: int,
    source_col: int,
    allergen_col: int,
    name_col: int,
    portion_blocks: list[int],
    warnings: list[str],
) -> list[DailyMenuItemPayload]:
    items: list[DailyMenuItemPayload] = []
    position = 1

    for row_number in range(start_row, end_row):
        source_text = _cell_text(sheet.cell(row_number, source_col).value)
        name = _cell_text(sheet.cell(row_number, name_col).value)

        if not source_text and not name:
            continue
        if _is_total_row(source_text) or _is_total_row(name):
            break
        if not name:
            warnings.append(f"Row {row_number}: skipped item without dish name.")
            continue

        kind = MenuItemKind.PRODUCT if _is_product_source(source_text) else MenuItemKind.DISH_CARD
        item = DailyMenuItemPayload(
            position=position,
            kind=kind,
            source_text=source_text or None,
            recipe_card_number=_extract_recipe_card_number(source_text)
            if kind == MenuItemKind.DISH_CARD
            else None,
            product_name_snapshot=name if kind == MenuItemKind.PRODUCT else None,
            name=name,
            allergen_codes=_parse_allergen_codes(sheet.cell(row_number, allergen_col).value),
            portions=_parse_portions(sheet, row_number, portion_blocks, warnings),
        )
        items.append(item)
        position += 1

    return items


def _parse_portions(
    sheet,
    row_number: int,
    portion_blocks: list[int],
    warnings: list[str],
) -> list[MenuPortionPayload]:
    portions: list[MenuPortionPayload] = []

    for age_group, start_col in zip(AGE_GROUPS_BY_BLOCK, portion_blocks, strict=True):
        yield_amount = _format_yield_amount(sheet.cell(row_number, start_col).value)
        if not yield_amount:
            warnings.append(f"Row {row_number}: missing yield amount for {age_group.value}.")
            continue

        portions.append(
            MenuPortionPayload(
                age_group=age_group,
                yield_amount=yield_amount,
                nutrition={
                    "kcal": _decimal_or_none(sheet.cell(row_number, start_col + 1).value),
                    "proteins": _decimal_or_none(sheet.cell(row_number, start_col + 2).value),
                    "fats": _decimal_or_none(sheet.cell(row_number, start_col + 3).value),
                    "carbs": _decimal_or_none(sheet.cell(row_number, start_col + 4).value),
                },
            )
        )

    return portions


def _parse_allergen_codes(value: Any) -> list[str]:
    text = _cell_text(value)
    if not text:
        return []
    codes: list[str] = []
    for raw_code in re.split(r"[,/;]", text):
        code = raw_code.strip().upper()
        if not code or code == "-":
            continue
        codes.append(code)
    return sorted(set(codes))


def _extract_recipe_card_number(source_text: str) -> str | None:
    match = re.search(r"ТК\s*№\s*([0-9]+(?:[._][0-9]+)*)", source_text, re.IGNORECASE)
    if match:
        return match.group(1).replace("_", ".").strip()

    match = re.search(r"№\s*([0-9]+(?:[._][0-9]+)*)", source_text, re.IGNORECASE)
    if match:
        return match.group(1).replace("_", ".").strip()

    return None


def _detect_cycle_week(sheet, sheet_name: str) -> int | None:
    for row_number in range(1, min(sheet.max_row, 5) + 1):
        for column in range(1, min(sheet.max_column, 6) + 1):
            text = _cell_text(sheet.cell(row_number, column).value).lower()
            match = re.search(r"(\d+)\s*-\s*й\s+тиждень", text)
            if match:
                return int(match.group(1))

    roman = sheet_name.strip().lower().replace(" ", "")
    if roman.startswith("іv") or roman.startswith("iv"):
        return 4
    if roman.startswith("ііі") or roman.startswith("iii"):
        return 3
    if roman.startswith("іі") or roman.startswith("ii"):
        return 2
    if roman.startswith("і") or roman.startswith("i"):
        return 1
    return None


def _default_import_title(*, filename: str, sheet_name: str, meal_type: MealType) -> str:
    meal_label = "Сніданки" if meal_type == MealType.BREAKFAST else "Обіди"
    file_label = filename.rsplit(".", 1)[0] if filename else "Імпорт меню"
    return f"{file_label}: {meal_label}, {sheet_name.strip()}"


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _is_total_row(value: str) -> bool:
    return value.strip().lower().startswith("всього")


def _is_product_source(source_text: str) -> bool:
    return "пром" in source_text.lower() and "вироб" in source_text.lower()


def _format_yield_amount(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _decimal_or_none(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    return Decimal(str(value))
