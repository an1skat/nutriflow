import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any

from app.core.config import get_settings
from app.modules.identity.models import AgeGroup
from app.modules.menus.models import (
    MealType,
    MenuImportDiagnostic,
    MenuImportDiagnosticLevel,
    MenuItemKind,
    Weekday,
    WeeklyMenu,
)
from app.modules.menus.schemas import (
    CreateWeeklyMenuRequest,
    DailyMenuItemPayload,
    DailyMenuPayload,
    MenuNutritionPayload,
    MenuPortionPayload,
)


class XlsxMenuError(ValueError):
    """The XLSX workbook cannot be parsed or generated safely."""


AGE_GROUPS_BY_BLOCK = [
    AgeGroup.SIX_TO_ELEVEN,
    AgeGroup.ELEVEN_TO_FOURTEEN,
    AgeGroup.FOURTEEN_TO_EIGHTEEN,
]

VISIBLE_MAX_COLUMN = 18
HEADER_GROUP_ROW = 1
HEADER_SUBHEADER_ROW = 2
WEEK_TITLE_ROW = 3
FIRST_CONTENT_ROW = 5
TEMPLATE_EMPTY_ITEM_ROWS = 6
SCHEMA_VERSION = 1
META_SHEET_NAME = "_meta"
TEMPLATE_SHEET_TITLES = [
    "І тиждень",
    "ІІ тиждень",
    "ІІІ тиждень",
    "ІV тиждень",
]

BASE_HEADERS = [
    "Збірник рецептур, № розкладки",
    "Алергени",
    "Найменування страв",
]
AGE_GROUP_HEADERS = [
    "Енергетична цінність для дітей 6-11 р.",
    "Енергетична цінність для дітей 11-14 р.",
    "Енергетична цінність для дітей 14-18 р.",
]
PORTION_SUBHEADERS = [
    "Вихід, г",
    "Енерго-цінність, ккал",
    "Білки, г",
    "Жири, г",
    "Вуглеводи, г",
]
OFFICIAL_PORTION_BLOCK_STARTS = [4, 9, 14]
DAY_TOTAL_LABEL = "Всього"

DAY_LABELS = {
    Weekday.MONDAY: "Понеділок",
    Weekday.TUESDAY: "Вівторок",
    Weekday.WEDNESDAY: "Середа",
    Weekday.THURSDAY: "Четвер",
    Weekday.FRIDAY: "П'ятниця",
    Weekday.SATURDAY: "Субота",
    Weekday.SUNDAY: "Неділя",
}

WEEKDAY_LOOKUP = {
    "понеділок": Weekday.MONDAY,
    "понедельник": Weekday.MONDAY,
    "пн": Weekday.MONDAY,
    "вівторок": Weekday.TUESDAY,
    "вторник": Weekday.TUESDAY,
    "вт": Weekday.TUESDAY,
    "середа": Weekday.WEDNESDAY,
    "среда": Weekday.WEDNESDAY,
    "ср": Weekday.WEDNESDAY,
    "четвер": Weekday.THURSDAY,
    "четверг": Weekday.THURSDAY,
    "чт": Weekday.THURSDAY,
    "п'ятниця": Weekday.FRIDAY,
    "п’ятниця": Weekday.FRIDAY,
    "пятниця": Weekday.FRIDAY,
    "пятница": Weekday.FRIDAY,
    "пт": Weekday.FRIDAY,
    "субота": Weekday.SATURDAY,
    "суббота": Weekday.SATURDAY,
    "сб": Weekday.SATURDAY,
    "неділя": Weekday.SUNDAY,
    "воскресенье": Weekday.SUNDAY,
    "нд": Weekday.SUNDAY,
}


@dataclass(slots=True)
class ParsedWeeklyMenuPreviewItem:
    sheet_name: str
    menu: CreateWeeklyMenuRequest
    item_rows: dict[tuple[str, int], int] = field(default_factory=dict)


@dataclass(slots=True)
class ParsedWeeklyMenuPreview:
    available_sheet_names: list[str]
    selected_sheet_name: str | None
    parsed_sheet_names: list[str]
    diagnostics: list[MenuImportDiagnostic]
    menus: list[ParsedWeeklyMenuPreviewItem]

    @property
    def menu(self) -> CreateWeeklyMenuRequest | None:
        return self.menus[0].menu if self.menus else None

    @property
    def commit_ready(self) -> bool:
        return bool(self.menus) and not any(
            diagnostic.level == MenuImportDiagnosticLevel.ERROR
            for diagnostic in self.diagnostics
        )


def preview_weekly_menu_workbook(
    content: bytes,
    *,
    filename: str,
    meal_type: MealType,
    sheet_name: str | None = None,
    title: str | None = None,
) -> ParsedWeeklyMenuPreview:
    openpyxl = _import_openpyxl()
    settings = get_settings()

    if not content:
        raise XlsxMenuError("Workbook is empty")
    if len(content) > settings.menu_import_max_file_bytes:
        raise XlsxMenuError(
            f"Workbook exceeds the {settings.menu_import_max_file_bytes} byte limit"
        )
    if not content.startswith(b"PK"):
        raise XlsxMenuError("File is not a valid .xlsx workbook")

    try:
        workbook = openpyxl.load_workbook(
            BytesIO(content),
            data_only=True,
            read_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise XlsxMenuError("Could not read the .xlsx workbook") from exc

    menu_sheet_names = [
        name for name in workbook.sheetnames if not name.strip().lower().startswith("_")
    ]
    if not menu_sheet_names:
        raise XlsxMenuError("Workbook does not contain menu worksheets")
    if len(menu_sheet_names) > settings.menu_import_max_sheet_count:
        raise XlsxMenuError(
            f"Workbook exceeds the {settings.menu_import_max_sheet_count} sheet limit"
        )
    if sheet_name is not None and sheet_name not in menu_sheet_names:
        raise XlsxMenuError("Requested worksheet was not found in workbook")

    target_sheet_names = [sheet_name] if sheet_name is not None else menu_sheet_names
    diagnostics: list[MenuImportDiagnostic] = []
    menus: list[ParsedWeeklyMenuPreviewItem] = []
    parsed_sheet_names: list[str] = []

    for target_sheet_name in target_sheet_names:
        sheet = workbook[target_sheet_name]
        last_row = _last_meaningful_row(sheet, settings.menu_import_max_rows_per_sheet)
        if last_row == 0:
            continue

        parsed = _parse_menu_sheet(
            sheet,
            last_row=last_row,
            filename=filename,
            meal_type=meal_type,
            title=title,
            multiple_sheets=len(target_sheet_names) > 1,
            diagnostics=diagnostics,
        )
        if parsed is None:
            continue
        menus.append(parsed)
        parsed_sheet_names.append(target_sheet_name)

    if not menus and not diagnostics:
        diagnostics.append(
            _diagnostic(
                level=MenuImportDiagnosticLevel.ERROR,
                code="no_menu_items",
                message="Workbook does not contain menu items",
                sheet_name=target_sheet_names[0],
            )
        )

    return ParsedWeeklyMenuPreview(
        available_sheet_names=menu_sheet_names,
        selected_sheet_name=(
            sheet_name
            or (menu_sheet_names[0] if len(target_sheet_names) == 1 else None)
        ),
        parsed_sheet_names=parsed_sheet_names,
        diagnostics=diagnostics,
        menus=menus,
    )


def build_template_workbook(*, weeks: int = 4) -> bytes:
    if weeks < 1 or weeks > 4:
        raise XlsxMenuError("Template can contain from one to four weeks")

    openpyxl = _import_openpyxl()
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    sheet_names = TEMPLATE_SHEET_TITLES[:weeks]
    for cycle_week, sheet_name in enumerate(sheet_names, start=1):
        sheet = workbook.create_sheet(sheet_name)
        _write_sheet_headers(sheet, cycle_week=cycle_week)
        row_number = FIRST_CONTENT_ROW
        for weekday in _school_weekdays():
            _write_day_separator(sheet, row_number, weekday)
            first_item_row = row_number + 1
            total_row = first_item_row + TEMPLATE_EMPTY_ITEM_ROWS
            _write_total_row(sheet, total_row, first_item_row, total_row - 1)
            row_number = total_row + 1
        _apply_print_settings(sheet, row_number - 1)

    _write_meta_sheet(workbook, [])
    return _save_workbook(workbook)


def build_export_workbook(menus: list[WeeklyMenu]) -> bytes:
    if not menus:
        raise XlsxMenuError("At least one weekly menu is required for export")
    if len(menus) > 4:
        raise XlsxMenuError("Can export at most four weekly menus in one workbook")

    openpyxl = _import_openpyxl()
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    used_titles: set[str] = set()
    for index, menu in enumerate(menus, start=1):
        cycle_week = menu.cycle_week or index
        preferred = _week_sheet_title(cycle_week)
        sheet_title = _unique_sheet_title(preferred, used_titles)
        used_titles.add(sheet_title)
        sheet = workbook.create_sheet(sheet_title)
        _write_sheet_headers(sheet, cycle_week=cycle_week)
        _write_menu_rows(sheet, menu)

    _write_meta_sheet(workbook, menus)
    return _save_workbook(workbook)


def _parse_menu_sheet(
    sheet,
    *,
    last_row: int,
    filename: str,
    meal_type: MealType,
    title: str | None,
    multiple_sheets: bool,
    diagnostics: list[MenuImportDiagnostic],
) -> ParsedWeeklyMenuPreviewItem | None:
    columns = _detect_columns(sheet, diagnostics)
    if columns is None:
        return None

    source_col, allergen_col, name_col, portion_starts = columns
    day_rows = _detect_day_rows(sheet, last_row=last_row, name_col=name_col)
    if not day_rows:
        diagnostics.append(
            _diagnostic(
                level=MenuImportDiagnosticLevel.ERROR,
                code="missing_weekdays",
                message="Worksheet does not contain weekday rows",
                sheet_name=sheet.title,
                row_number=FIRST_CONTENT_ROW,
                column_number=name_col,
            )
        )
        return None

    days: list[DailyMenuPayload] = []
    item_rows: dict[tuple[str, int], int] = {}
    seen_weekdays: set[Weekday] = set()

    for index, (day_row, weekday) in enumerate(day_rows):
        if weekday in seen_weekdays:
            diagnostics.append(
                _diagnostic(
                    level=MenuImportDiagnosticLevel.ERROR,
                    code="duplicate_weekday",
                    message=f"Weekday {DAY_LABELS[weekday]} appears more than once",
                    sheet_name=sheet.title,
                    row_number=day_row,
                    column_number=name_col,
                )
            )
            continue
        seen_weekdays.add(weekday)

        next_day_row = day_rows[index + 1][0] if index + 1 < len(day_rows) else last_row + 1
        items, rows_by_position = _parse_day_items(
            sheet,
            start_row=day_row + 1,
            end_row=next_day_row,
            source_col=source_col,
            allergen_col=allergen_col,
            name_col=name_col,
            portion_starts=portion_starts,
            sheet_name=sheet.title,
            diagnostics=diagnostics,
        )
        if items:
            days.append(DailyMenuPayload(weekday=weekday, items=items))
            item_rows.update(
                {(weekday.value, position): row for position, row in rows_by_position.items()}
            )

    if not days:
        diagnostics.append(
            _diagnostic(
                level=MenuImportDiagnosticLevel.ERROR,
                code="no_menu_items",
                message="Worksheet does not contain menu items",
                sheet_name=sheet.title,
            )
        )
        return None

    cycle_week = _detect_cycle_week(sheet, sheet.title)
    resolved_title = _resolved_import_title(
        title=title,
        filename=filename,
        sheet_name=sheet.title,
        meal_type=meal_type,
        multiple_sheets=multiple_sheets,
    )
    try:
        menu = CreateWeeklyMenuRequest(
            title=resolved_title,
            meal_type=meal_type,
            cycle_week=cycle_week,
            days=days,
            source_file_name=filename,
            source_sheet_name=sheet.title,
        )
    except ValueError as exc:
        diagnostics.append(
            _diagnostic(
                level=MenuImportDiagnosticLevel.ERROR,
                code="invalid_menu",
                message=str(exc),
                sheet_name=sheet.title,
            )
        )
        return None

    return ParsedWeeklyMenuPreviewItem(
        sheet_name=sheet.title,
        menu=menu,
        item_rows=item_rows,
    )


def _detect_columns(
    sheet,
    diagnostics: list[MenuImportDiagnostic],
) -> tuple[int, int, int, list[int]] | None:
    source_col: int | None = None
    allergen_col: int | None = None
    name_col: int | None = None

    scan_max_column = min(max(sheet.max_column, VISIBLE_MAX_COLUMN), 64)
    for row_number in range(1, 5):
        for column_number in range(1, scan_max_column + 1):
            text = _normalize_text(sheet.cell(row_number, column_number).value)
            if not text:
                continue
            if source_col is None and (
                "збірник" in text or "розкладки" in text or "техкарт" in text
            ):
                source_col = column_number
            elif allergen_col is None and "алерген" in text:
                allergen_col = column_number
            elif name_col is None and (
                "найменування" in text or "назва страв" in text
            ):
                name_col = column_number

    if source_col is None or allergen_col is None or name_col is None:
        diagnostics.append(
            _diagnostic(
                level=MenuImportDiagnosticLevel.ERROR,
                code="invalid_headers",
                message=(
                    "Could not detect recipe-card, allergen, and dish-name columns "
                    "in the first four rows"
                ),
                sheet_name=sheet.title,
                row_number=1,
            )
        )
        return None

    portion_starts: list[int] = []
    for row_number in range(1, 5):
        for column_number in range(name_col + 1, scan_max_column + 1):
            text = _normalize_text(sheet.cell(row_number, column_number).value)
            if text.startswith("вихід") and column_number not in portion_starts:
                portion_starts.append(column_number)
    portion_starts.sort()

    if len(portion_starts) < 3:
        diagnostics.append(
            _diagnostic(
                level=MenuImportDiagnosticLevel.ERROR,
                code="invalid_age_group_headers",
                message="Could not detect all three age-group nutrition blocks",
                sheet_name=sheet.title,
                row_number=HEADER_SUBHEADER_ROW,
                column_number=name_col + 1,
            )
        )
        return None

    portion_starts = portion_starts[:3]
    if any(
        next_start - start < 5
        for start, next_start in zip(portion_starts, portion_starts[1:], strict=False)
    ):
        diagnostics.append(
            _diagnostic(
                level=MenuImportDiagnosticLevel.ERROR,
                code="invalid_age_group_layout",
                message="Every age-group block must contain five consecutive columns",
                sheet_name=sheet.title,
                row_number=HEADER_SUBHEADER_ROW,
                column_number=portion_starts[0],
            )
        )
        return None

    return source_col, allergen_col, name_col, portion_starts


def _detect_day_rows(
    sheet,
    *,
    last_row: int,
    name_col: int,
) -> list[tuple[int, Weekday]]:
    day_rows: list[tuple[int, Weekday]] = []
    candidate_columns = sorted({name_col, *range(max(1, name_col - 2), name_col + 1)})

    for row_number in range(1, last_row + 1):
        for column_number in candidate_columns:
            weekday = _weekday_from_value(sheet.cell(row_number, column_number).value)
            if weekday is not None:
                day_rows.append((row_number, weekday))
                break

    return day_rows


def _parse_day_items(
    sheet,
    *,
    start_row: int,
    end_row: int,
    source_col: int,
    allergen_col: int,
    name_col: int,
    portion_starts: list[int],
    sheet_name: str,
    diagnostics: list[MenuImportDiagnostic],
) -> tuple[list[DailyMenuItemPayload], dict[int, int]]:
    items: list[DailyMenuItemPayload] = []
    rows_by_position: dict[int, int] = {}

    for row_number in range(start_row, end_row):
        if _row_contains_total(sheet, row_number, source_col, allergen_col, name_col):
            break

        source_text = _cell_text(sheet.cell(row_number, source_col).value)
        name = _cell_text(sheet.cell(row_number, name_col).value)
        allergen_value = sheet.cell(row_number, allergen_col).value
        row_has_nutrition = any(
            _cell_text(sheet.cell(row_number, column_number).value)
            for start_col in portion_starts
            for column_number in range(start_col, start_col + 5)
        )

        if not source_text and not name and not row_has_nutrition:
            continue
        if not name:
            diagnostics.append(
                _diagnostic(
                    level=MenuImportDiagnosticLevel.ERROR,
                    code="missing_name",
                    message="Menu row is missing a dish or product name",
                    sheet_name=sheet_name,
                    row_number=row_number,
                    column_number=name_col,
                )
            )
            continue

        portions = _parse_portions(
            sheet,
            row_number=row_number,
            portion_starts=portion_starts,
            sheet_name=sheet_name,
            diagnostics=diagnostics,
        )
        if len(portions) != len(AGE_GROUPS_BY_BLOCK):
            diagnostics.append(
                _diagnostic(
                    level=MenuImportDiagnosticLevel.ERROR,
                    code="incomplete_age_groups",
                    message="Menu row must include all three age-group blocks",
                    sheet_name=sheet_name,
                    row_number=row_number,
                    column_number=portion_starts[0],
                )
            )
            continue

        kind = MenuItemKind.PRODUCT if _is_product_source(source_text) else MenuItemKind.DISH_CARD
        position = len(items) + 1
        try:
            item = DailyMenuItemPayload(
                position=position,
                kind=kind,
                source_text=source_text or None,
                recipe_card_number=(
                    None
                    if kind == MenuItemKind.PRODUCT
                    else _extract_recipe_card_number(source_text)
                ),
                product_name_snapshot=name if kind == MenuItemKind.PRODUCT else None,
                name=name,
                allergen_codes=_parse_allergen_codes(allergen_value),
                portions=portions,
            )
        except ValueError as exc:
            diagnostics.append(
                _diagnostic(
                    level=MenuImportDiagnosticLevel.ERROR,
                    code="invalid_item",
                    message=str(exc),
                    sheet_name=sheet_name,
                    row_number=row_number,
                    column_number=name_col,
                )
            )
            continue

        items.append(item)
        rows_by_position[position] = row_number

    return items, rows_by_position


def _parse_portions(
    sheet,
    *,
    row_number: int,
    portion_starts: list[int],
    sheet_name: str,
    diagnostics: list[MenuImportDiagnostic],
) -> list[MenuPortionPayload]:
    portions: list[MenuPortionPayload] = []

    for age_group, start_col in zip(
        AGE_GROUPS_BY_BLOCK,
        portion_starts,
        strict=True,
    ):
        yield_amount = _format_yield_amount(sheet.cell(row_number, start_col).value)
        raw_nutrition = [
            sheet.cell(row_number, start_col + offset).value for offset in range(1, 5)
        ]
        if not yield_amount and not any(_cell_text(value) for value in raw_nutrition):
            continue
        if not yield_amount:
            diagnostics.append(
                _diagnostic(
                    level=MenuImportDiagnosticLevel.ERROR,
                    code="missing_yield",
                    message=f"{age_group.value} portion is missing yield",
                    sheet_name=sheet_name,
                    row_number=row_number,
                    column_number=start_col,
                )
            )
            continue

        nutrition = MenuNutritionPayload(
            kcal=_decimal_or_diagnostic(
                raw_nutrition[0],
                sheet_name=sheet_name,
                row_number=row_number,
                column_number=start_col + 1,
                diagnostics=diagnostics,
            ),
            proteins=_decimal_or_diagnostic(
                raw_nutrition[1],
                sheet_name=sheet_name,
                row_number=row_number,
                column_number=start_col + 2,
                diagnostics=diagnostics,
            ),
            fats=_decimal_or_diagnostic(
                raw_nutrition[2],
                sheet_name=sheet_name,
                row_number=row_number,
                column_number=start_col + 3,
                diagnostics=diagnostics,
            ),
            carbs=_decimal_or_diagnostic(
                raw_nutrition[3],
                sheet_name=sheet_name,
                row_number=row_number,
                column_number=start_col + 4,
                diagnostics=diagnostics,
            ),
        )
        portions.append(
            MenuPortionPayload(
                age_group=age_group,
                yield_amount=yield_amount,
                nutrition=nutrition,
            )
        )

    return portions


def _write_sheet_headers(sheet, *, cycle_week: int) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    green = "166534"
    pale_green = "DCFCE7"
    light_green = "F0FDF4"
    border_side = Side(style="thin", color="94A3B8")
    border = Border(
        left=border_side,
        right=border_side,
        top=border_side,
        bottom=border_side,
    )

    sheet.merge_cells("A1:A2")
    sheet.merge_cells("B1:B2")
    sheet.merge_cells("C1:C2")
    sheet.merge_cells("D1:H1")
    sheet.merge_cells("I1:M1")
    sheet.merge_cells("N1:R1")

    for column_number, header in enumerate(BASE_HEADERS, start=1):
        sheet.cell(HEADER_GROUP_ROW, column_number, header)
    for start_col, header in zip(
        OFFICIAL_PORTION_BLOCK_STARTS,
        AGE_GROUP_HEADERS,
        strict=True,
    ):
        sheet.cell(HEADER_GROUP_ROW, start_col, header)
        for offset, subheader in enumerate(PORTION_SUBHEADERS):
            sheet.cell(HEADER_SUBHEADER_ROW, start_col + offset, subheader)

    for row in sheet.iter_rows(
        min_row=HEADER_GROUP_ROW,
        max_row=HEADER_SUBHEADER_ROW,
        min_col=1,
        max_col=VISIBLE_MAX_COLUMN,
    ):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.font = Font(bold=True, color="FFFFFF" if cell.row == 1 else "14532D")
            cell.fill = PatternFill(
                "solid",
                fgColor=green if cell.row == 1 else light_green,
            )

    sheet.merge_cells(
        start_row=WEEK_TITLE_ROW,
        start_column=1,
        end_row=WEEK_TITLE_ROW,
        end_column=18,
    )
    week_cell = sheet.cell(WEEK_TITLE_ROW, 1, f"{cycle_week}-й тиждень")
    week_cell.font = Font(bold=True, size=13, color="14532D")
    week_cell.fill = PatternFill("solid", fgColor=pale_green)
    week_cell.alignment = Alignment(horizontal="center", vertical="center")
    week_cell.border = border

    widths = [30, 16, 42, 12, 16, 11, 11, 14, 12, 16, 11, 11, 14, 12, 16, 11, 11, 14]
    for column_number, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column_number)].width = width

    sheet.row_dimensions[1].height = 42
    sheet.row_dimensions[2].height = 50
    sheet.row_dimensions[3].height = 24
    sheet.freeze_panes = "D5"


def _write_day_separator(sheet, row_number: int, weekday: Weekday) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    side = Side(style="thin", color="94A3B8")
    for column_number in range(1, VISIBLE_MAX_COLUMN + 1):
        cell = sheet.cell(row_number, column_number)
        cell.fill = PatternFill("solid", fgColor="DCFCE7")
        cell.border = Border(left=side, right=side, top=side, bottom=side)
    cell = sheet.cell(row_number, 3, DAY_LABELS[weekday])
    cell.font = Font(bold=True, color="14532D")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[row_number].height = 24


def _write_total_row(
    sheet,
    row_number: int,
    first_item_row: int,
    last_item_row: int,
) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    side = Side(style="thin", color="94A3B8")
    for column_number in range(1, VISIBLE_MAX_COLUMN + 1):
        cell = sheet.cell(row_number, column_number)
        cell.fill = PatternFill("solid", fgColor="FEF3C7")
        cell.border = Border(left=side, right=side, top=side, bottom=side)
        cell.font = Font(bold=True, color="713F12")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.cell(row_number, 3, DAY_TOTAL_LABEL)
    for column_number in range(4, VISIBLE_MAX_COLUMN + 1):
        column_letter = get_column_letter(column_number)
        sheet.cell(
            row_number,
            column_number,
            f"=SUM({column_letter}{first_item_row}:{column_letter}{last_item_row})",
        )
        sheet.cell(row_number, column_number).number_format = "0.00"


def _write_menu_rows(sheet, menu: WeeklyMenu) -> None:
    from openpyxl.styles import Alignment, Border, Side

    row_number = FIRST_CONTENT_ROW
    side = Side(style="thin", color="CBD5E1")
    border = Border(left=side, right=side, top=side, bottom=side)

    for day in sorted(menu.days, key=lambda value: list(Weekday).index(value.weekday)):
        _write_day_separator(sheet, row_number, day.weekday)
        first_item_row = row_number + 1
        row_number = first_item_row

        for item in sorted(day.items, key=lambda value: value.position):
            sheet.cell(row_number, 1, _export_source_text(item))
            sheet.cell(row_number, 2, ", ".join(item.allergen_codes))
            sheet.cell(row_number, 3, item.name)

            portions_by_age = {portion.age_group: portion for portion in item.portions}
            for age_group, start_col in zip(
                AGE_GROUPS_BY_BLOCK,
                OFFICIAL_PORTION_BLOCK_STARTS,
                strict=True,
            ):
                portion = portions_by_age.get(age_group)
                if portion is None:
                    continue
                sheet.cell(row_number, start_col, portion.yield_amount)
                sheet.cell(row_number, start_col + 1, _decimal_for_excel(portion.nutrition.kcal))
                sheet.cell(
                    row_number,
                    start_col + 2,
                    _decimal_for_excel(portion.nutrition.proteins),
                )
                sheet.cell(row_number, start_col + 3, _decimal_for_excel(portion.nutrition.fats))
                sheet.cell(row_number, start_col + 4, _decimal_for_excel(portion.nutrition.carbs))

            for column_number in range(1, VISIBLE_MAX_COLUMN + 1):
                cell = sheet.cell(row_number, column_number)
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            for column_number in range(5, VISIBLE_MAX_COLUMN + 1):
                if column_number not in OFFICIAL_PORTION_BLOCK_STARTS:
                    sheet.cell(row_number, column_number).number_format = "0.00"
            sheet.row_dimensions[row_number].height = 34
            row_number += 1

        _write_total_row(sheet, row_number, first_item_row, row_number - 1)
        row_number += 1

    _apply_print_settings(sheet, row_number - 1)


def _write_meta_sheet(workbook, menus: list[WeeklyMenu]) -> None:
    meta = workbook.create_sheet(META_SHEET_NAME)
    meta.append(["key", "value"])
    meta.append(["schema_version", SCHEMA_VERSION])
    meta.append(["document_type", "weekly_menu_batch"])
    meta.append(["generated_at", datetime.now(UTC).isoformat()])
    meta.append([])
    meta.append(
        [
            "sheet_name",
            "cycle_week",
            "menu_id",
            "school_id",
            "source_menu_id",
            "meal_type",
        ]
    )
    for index, menu in enumerate(menus, start=1):
        meta.append(
            [
                workbook.worksheets[index - 1].title,
                menu.cycle_week,
                str(menu.id or ""),
                str(menu.school_id or ""),
                str(menu.source_menu_id or ""),
                menu.meal_type.value,
            ]
        )
    meta.sheet_state = "hidden"


def _apply_print_settings(sheet, last_row: int) -> None:
    sheet.auto_filter.ref = f"A1:R{max(last_row, 3)}"
    sheet.print_area = f"A1:R{max(last_row, 3)}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"
    sheet.page_margins.left = 0.2
    sheet.page_margins.right = 0.2
    sheet.page_margins.top = 0.3
    sheet.page_margins.bottom = 0.3


def _last_meaningful_row(sheet, max_rows: int) -> int:
    scan_columns = min(max(sheet.max_column, VISIBLE_MAX_COLUMN), 64)
    scan_to = min(sheet.max_row, max_rows)
    last_row = 0

    for row_number in range(1, scan_to + 1):
        if any(
            _cell_text(sheet.cell(row_number, column_number).value)
            for column_number in range(1, scan_columns + 1)
        ):
            last_row = row_number

    if sheet.max_row > max_rows:
        overflow_scan_to = min(sheet.max_row, max_rows + 200)
        for row_number in range(max_rows + 1, overflow_scan_to + 1):
            if any(
                _cell_text(sheet.cell(row_number, column_number).value)
                for column_number in range(1, scan_columns + 1)
            ):
                raise XlsxMenuError(
                    f"Worksheet {sheet.title!r} exceeds the {max_rows} meaningful-row limit"
                )

    return last_row


def _weekday_from_value(value: Any) -> Weekday | None:
    text = _normalize_text(value).strip(" :.-")
    return WEEKDAY_LOOKUP.get(text)


def _row_contains_total(
    sheet,
    row_number: int,
    source_col: int,
    allergen_col: int,
    name_col: int,
) -> bool:
    for column_number in {source_col, allergen_col, name_col}:
        if _normalize_text(sheet.cell(row_number, column_number).value).startswith("всього"):
            return True
    return False


def _parse_allergen_codes(value: Any) -> list[str]:
    text = _cell_text(value)
    if not text:
        return []
    return sorted(
        {
            code
            for raw_code in re.split(r"[,/;]", text)
            if (code := raw_code.strip().upper()) and code != "-"
        }
    )


def _extract_recipe_card_number(source_text: str) -> str | None:
    if not source_text:
        return None
    matches = re.findall(r"(?:ТК\s*№\s*|№\s*)(\d+(?:[._/-]\d+)*)", source_text, re.IGNORECASE)
    if matches:
        return matches[-1].replace("_", ".").strip()
    match = re.search(r"(\d+(?:[./-]\d+)+)", source_text)
    return match.group(1) if match else None


def _is_product_source(source_text: str) -> bool:
    normalized = _normalize_text(source_text)
    return "пром" in normalized and "вироб" in normalized


def _decimal_or_diagnostic(
    value: Any,
    *,
    sheet_name: str,
    row_number: int,
    column_number: int,
    diagnostics: list[MenuImportDiagnostic],
) -> Decimal | None:
    if value in (None, ""):
        return None
    if isinstance(value, str) and value.startswith("="):
        diagnostics.append(
            _diagnostic(
                level=MenuImportDiagnosticLevel.ERROR,
                code="formula_not_allowed",
                message="Formula cells are not allowed in menu-item nutrition rows",
                sheet_name=sheet_name,
                row_number=row_number,
                column_number=column_number,
            )
        )
        return None

    try:
        return Decimal(str(value).strip().replace(" ", "").replace(",", "."))
    except (ArithmeticError, InvalidOperation, ValueError):
        diagnostics.append(
            _diagnostic(
                level=MenuImportDiagnosticLevel.ERROR,
                code="invalid_numeric_value",
                message="Expected a numeric nutrition value",
                sheet_name=sheet_name,
                row_number=row_number,
                column_number=column_number,
            )
        )
        return None


def _diagnostic(
    *,
    level: MenuImportDiagnosticLevel,
    code: str,
    message: str,
    sheet_name: str,
    row_number: int | None = None,
    column_number: int | None = None,
) -> MenuImportDiagnostic:
    column_letter = _column_letter(column_number) if column_number is not None else None
    cell = (
        f"{column_letter}{row_number}"
        if column_letter is not None and row_number is not None
        else None
    )
    return MenuImportDiagnostic(
        level=level,
        code=code,
        message=message,
        sheet_name=sheet_name,
        row_number=row_number,
        column_letter=column_letter,
        cell=cell,
    )


def _detect_cycle_week(sheet, sheet_name: str) -> int | None:
    for row_number in range(1, min(sheet.max_row, 5) + 1):
        for column_number in range(1, min(sheet.max_column, 18) + 1):
            text = _normalize_text(sheet.cell(row_number, column_number).value)
            match = re.search(r"(\d+)\s*-?\s*(?:й|тиждень)", text)
            if match:
                return int(match.group(1))

    normalized_name = _normalize_text(sheet_name).replace(" ", "")
    roman_match = re.match(r"^(iv|iii|ii|i|іv|ііі|іі|і)", normalized_name)
    if roman_match:
        return {
            "i": 1,
            "і": 1,
            "ii": 2,
            "іі": 2,
            "iii": 3,
            "ііі": 3,
            "iv": 4,
            "іv": 4,
        }[roman_match.group(1)]

    number_match = re.search(r"\d+", normalized_name)
    return int(number_match.group()) if number_match else None


def _resolved_import_title(
    *,
    title: str | None,
    filename: str,
    sheet_name: str,
    meal_type: MealType,
    multiple_sheets: bool,
) -> str:
    if title:
        return f"{title} - {sheet_name}" if multiple_sheets else title
    meal_label = "Сніданки" if meal_type == MealType.BREAKFAST else "Обіди"
    file_label = Path(filename or "menu").stem.replace("_", " ").strip() or "Меню"
    return f"{file_label}: {meal_label}, {sheet_name.strip()}"


def _export_source_text(item) -> str:
    if item.source_text:
        return item.source_text
    if item.kind == MenuItemKind.PRODUCT:
        return "пром. вироб."
    if item.recipe_card_number:
        return f"ТК № {item.recipe_card_number}"
    return ""


def _decimal_for_excel(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


def _format_yield_amount(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _week_sheet_title(cycle_week: int) -> str:
    if 1 <= cycle_week <= len(TEMPLATE_SHEET_TITLES):
        return TEMPLATE_SHEET_TITLES[cycle_week - 1]
    return f"Тиждень {cycle_week}"


def _unique_sheet_title(preferred: str, used_titles: set[str]) -> str:
    sanitized = re.sub(r'[:\\/?*\[\]]', "-", preferred).strip() or "Меню"
    sanitized = sanitized[:31]
    if sanitized not in used_titles:
        return sanitized
    suffix = 2
    while True:
        candidate = f"{sanitized[: 28 - len(str(suffix))]} ({suffix})"
        if candidate not in used_titles:
            return candidate
        suffix += 1


def _school_weekdays() -> list[Weekday]:
    return [
        Weekday.MONDAY,
        Weekday.TUESDAY,
        Weekday.WEDNESDAY,
        Weekday.THURSDAY,
        Weekday.FRIDAY,
    ]


def _normalize_text(value: Any) -> str:
    text = _cell_text(value).lower().replace("’", "'").replace("`", "'")
    return re.sub(r"\s+", " ", text)


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _column_letter(column_number: int) -> str:
    result = ""
    current = column_number
    while current > 0:
        current, remainder = divmod(current - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _save_workbook(workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _import_openpyxl():
    try:
        import openpyxl
    except ModuleNotFoundError as exc:
        raise XlsxMenuError("openpyxl is required for menu workbooks") from exc
    return openpyxl
