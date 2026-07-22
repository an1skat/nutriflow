from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.modules.menu_requirements.models import MenuRequirement
from app.modules.menu_requirements.schemas import (
    MenuRequirementReportGroupResponse,
    MenuRequirementReportResponse,
)

TITLE_FILL = PatternFill("solid", fgColor="166534")
HEADER_FILL = PatternFill("solid", fgColor="DCFCE7")
TOTAL_FILL = PatternFill("solid", fgColor="D1FAE5")
THIN_GRAY = Side(style="thin", color="CBD5E1")
TABLE_BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)
MEAL_TYPE_LABELS = {"breakfast": "Сніданок", "lunch": "Обід"}
GRANULARITY_LABELS = {"day": "день", "week": "тиждень", "month": "місяць"}


def build_menu_requirement_workbook(
    requirement: MenuRequirement,
    *,
    school_name: str,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _safe_sheet_title(requirement.school_group_name)

    metadata = [
        ("Школа", school_name),
        ("Вікова група", requirement.age_group.value),
        ("Дата", requirement.service_date.strftime("%d.%m.%Y")),
        ("Прийом їжі", MEAL_TYPE_LABELS[requirement.meal_type.value]),
        ("Меню", requirement.menu_title),
        ("Версія", requirement.revision),
    ]
    table_row = _write_sheet_heading(sheet, "МЕНЮ-ВИМОГА", metadata)
    headers = [
        "Інгредієнт",
        *[
            f"{dish.name}\nВихід: {dish.yield_amount} г\nДітей: {dish.children_count}"
            for dish in requirement.dishes
        ],
        "Разом на одну особу, г",
        "До видачі, г",
    ]
    _write_table_header(sheet, table_row, headers)

    dish_ids = [dish.menu_item_id for dish in requirement.dishes]
    for row_number, ingredient in enumerate(
        requirement.ingredient_rows,
        start=table_row + 1,
    ):
        cells_by_dish = {cell.menu_item_id: cell for cell in ingredient.cells}
        values: list[Any] = [ingredient.ingredient_name]
        values.extend(
            _excel_number(cells_by_dish[dish_id].net_per_person_g)
            if dish_id in cells_by_dish
            else None
            for dish_id in dish_ids
        )
        values.extend(
            [
                _excel_number(ingredient.per_person_total_g),
                ingredient.issue_total_rounded_g,
            ]
        )
        _write_table_row(sheet, row_number, values, total_columns=2)

    _finish_table_sheet(
        sheet,
        header_row=table_row,
        last_row=table_row + len(requirement.ingredient_rows),
        column_count=len(headers),
    )
    return _workbook_bytes(workbook)


def build_menu_requirement_report_workbook(
    report: MenuRequirementReportResponse,
) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)

    if not report.groups:
        sheet = workbook.create_sheet("Меню-вимога")
        metadata = _report_metadata(report)
        message_row = _write_sheet_heading(sheet, "МЕНЮ-ВИМОГА", metadata)
        sheet.cell(message_row, 1, "Немає даних меню-вимог за обраний період.")
        sheet.column_dimensions["A"].width = 48
    else:
        used_titles: set[str] = set()
        for group in report.groups:
            title = _unique_sheet_title(group.school_group_name, used_titles)
            sheet = workbook.create_sheet(title)
            _write_report_group(sheet, report, group)

    return _workbook_bytes(workbook)


def _write_report_group(
    sheet: Worksheet,
    report: MenuRequirementReportResponse,
    group: MenuRequirementReportGroupResponse,
) -> None:
    metadata = [
        *_report_metadata(report),
        ("Вікова група", group.age_group.value),
    ]
    table_row = _write_sheet_heading(sheet, "МЕНЮ-ВИМОГА", metadata)
    headers = [
        "Інгредієнт",
        *[
            (f"{dish.name}\nВихід: {dish.yield_amount} г\nДітей: {dish.children_count_total}")
            for dish in group.dishes
        ],
        "Разом нетто, г",
        "До видачі, г",
    ]
    _write_table_header(sheet, table_row, headers)

    dish_keys = [dish.aggregate_key for dish in group.dishes]
    for row_number, ingredient in enumerate(group.ingredient_rows, start=table_row + 1):
        cells_by_dish = {cell.dish_key: cell for cell in ingredient.cells}
        values: list[Any] = [ingredient.ingredient_name]
        values.extend(
            cells_by_dish[dish_key].issue_total_rounded_g if dish_key in cells_by_dish else None
            for dish_key in dish_keys
        )
        values.extend(
            [
                _excel_number(ingredient.per_person_total_g),
                ingredient.issue_total_rounded_g,
            ]
        )
        _write_table_row(sheet, row_number, values, total_columns=2)

    _finish_table_sheet(
        sheet,
        header_row=table_row,
        last_row=table_row + len(group.ingredient_rows),
        column_count=len(headers),
    )


def _report_metadata(report: MenuRequirementReportResponse) -> list[tuple[str, Any]]:
    meal_type = MEAL_TYPE_LABELS[report.meal_type.value] if report.meal_type else "Усі"
    return [
        ("Школа", report.school_name),
        (
            "Період",
            f"{report.date_from.strftime('%d.%m.%Y')} – {report.date_to.strftime('%d.%m.%Y')}",
        ),
        ("Групування", GRANULARITY_LABELS[report.granularity.value]),
        ("Прийом їжі", meal_type),
    ]


def _write_sheet_heading(
    sheet: Worksheet,
    title: str,
    metadata: list[tuple[str, Any]],
) -> int:
    sheet.cell(1, 1, title)
    sheet.cell(1, 1).font = Font(bold=True, color="FFFFFF", size=14)
    sheet.cell(1, 1).fill = TITLE_FILL
    sheet.cell(1, 1).alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 26

    for row_number, (label, value) in enumerate(metadata, start=2):
        sheet.cell(row_number, 1, label).font = Font(bold=True, color="475569")
        sheet.cell(row_number, 2, value)

    return len(metadata) + 3


def _write_table_header(sheet: Worksheet, row_number: int, headers: list[str]) -> None:
    for column_number, value in enumerate(headers, start=1):
        cell = sheet.cell(row_number, column_number, value)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True, color="14532D")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = TABLE_BORDER
    sheet.row_dimensions[row_number].height = 54


def _write_table_row(
    sheet: Worksheet,
    row_number: int,
    values: list[Any],
    *,
    total_columns: int,
) -> None:
    total_column_start = len(values) - total_columns + 1
    for column_number, value in enumerate(values, start=1):
        cell = sheet.cell(row_number, column_number, value)
        cell.border = TABLE_BORDER
        cell.alignment = Alignment(
            horizontal="left" if column_number == 1 else "right",
            vertical="center",
            wrap_text=column_number == 1,
        )
        if column_number >= total_column_start:
            cell.fill = TOTAL_FILL
            cell.font = Font(bold=True)
        if column_number > 1 and value is not None:
            cell.number_format = "0.######"


def _finish_table_sheet(
    sheet: Worksheet,
    *,
    header_row: int,
    last_row: int,
    column_count: int,
) -> None:
    sheet.freeze_panes = f"B{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(column_count)}{last_row}"
    sheet.column_dimensions["A"].width = 30
    for column_number in range(2, column_count - 1):
        sheet.column_dimensions[get_column_letter(column_number)].width = 22
    sheet.column_dimensions[get_column_letter(column_count - 1)].width = 22
    sheet.column_dimensions[get_column_letter(column_count)].width = 18
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"{header_row}:{header_row}"


def _excel_number(value: Any) -> float:
    return float(value)


def _safe_sheet_title(value: str) -> str:
    cleaned = "".join("-" if character in "[]:*?/\\" else character for character in value)
    return cleaned.strip()[:31] or "Меню-вимога"


def _unique_sheet_title(value: str, used_titles: set[str]) -> str:
    base = _safe_sheet_title(value)
    candidate = base
    index = 2
    while candidate.casefold() in used_titles:
        suffix = f" ({index})"
        candidate = f"{base[: 31 - len(suffix)]}{suffix}"
        index += 1
    used_titles.add(candidate.casefold())
    return candidate


def _workbook_bytes(workbook: Workbook) -> bytes:
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
