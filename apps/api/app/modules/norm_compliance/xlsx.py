from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.modules.norm_compliance.schemas import ComplianceStatus, NormComplianceReportResponse

STATUS_LABELS = {
    ComplianceStatus.COMPLETE: "В нормі",
    ComplianceStatus.UNDER: "Нижче норми",
    ComplianceStatus.OVER: "Вище норми",
    ComplianceStatus.MISSING: "Немає меню-вимоги",
    ComplianceStatus.STALE: "Застаріло",
    ComplianceStatus.UNMAPPED: "Не визначено групу продукту",
    ComplianceStatus.MIXED: "Змішаний статус",
}
STATUS_COLORS = {
    ComplianceStatus.COMPLETE: "ECFDF5",
    ComplianceStatus.UNDER: "FFF1F2",
    ComplianceStatus.OVER: "FFF1F2",
    ComplianceStatus.MISSING: "FFFBEB",
    ComplianceStatus.STALE: "FFFBEB",
    ComplianceStatus.UNMAPPED: "FFFBEB",
    ComplianceStatus.MIXED: "FFFBEB",
}
MEAL_TYPE_LABELS = {"breakfast": "Сніданок", "lunch": "Обід"}
AGE_GROUP_LABELS = {"6-11": "6-11 років", "11-14": "11-14 років", "14-18": "14-18 років"}
UNIT_LABELS = {"g": "г", "ml": "мл", "item": "шт.", "portion": "порц."}
HEADERS = [
    "Група продуктів",
    "Частота",
    "Норма порцій",
    "Факт порцій",
    "Норма нетто / обсяг",
    "Факт",
    "Відхилення",
    "Одиниця",
    "% виконання",
    "Статус",
]


def build_norm_compliance_workbook(report: NormComplianceReportResponse) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Дотримання норм"
    sheet.append(["ДОТРИМАННЯ НОРМ ХАРЧУВАННЯ"])
    sheet.append(["Школа", report.school_name])
    sheet.append(
        [
            "Період",
            f"{report.date_from.strftime('%d.%m.%Y')} - {report.date_to.strftime('%d.%m.%Y')}",
        ]
    )
    sheet.append(["Загальний статус", STATUS_LABELS[report.status]])
    sheet.append([])
    sheet.append(HEADERS)

    sheet.merge_cells("A1:J1")
    sheet["A1"].fill = PatternFill("solid", fgColor="166534")
    sheet["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 26
    for row_number in range(2, 5):
        sheet.cell(row_number, 1).font = Font(bold=True, color="475569")

    border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    for cell in sheet[6]:
        cell.fill = PatternFill("solid", fgColor="DCFCE7")
        cell.font = Font(bold=True, color="14532D")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    sheet.row_dimensions[6].height = 42

    for group in report.groups:
        for section in group.sections:
            section_row = sheet.max_row + 1
            sheet.cell(
                section_row,
                1,
                (
                    f"{group.school_group_name} · {AGE_GROUP_LABELS[group.age_group.value]} · "
                    f"{MEAL_TYPE_LABELS[section.meal_type.value]} · "
                    f"{len(section.expected_dates)} дн."
                ),
            )
            sheet.merge_cells(
                start_row=section_row,
                start_column=1,
                end_row=section_row,
                end_column=len(HEADERS),
            )
            sheet.cell(section_row, 1).fill = PatternFill("solid", fgColor="E2E8F0")
            sheet.cell(section_row, 1).font = Font(bold=True, color="334155")

            for row in section.rows:
                sheet.append(
                    [
                        row.normative_group_name,
                        row.frequency,
                        float(row.required_portions),
                        float(row.actual_portions),
                        float(row.required_amount),
                        float(row.actual_amount),
                        float(row.deviation),
                        UNIT_LABELS[row.unit.value],
                        float(row.percent) if row.percent is not None else None,
                        STATUS_LABELS[row.status],
                    ]
                )
                fill = PatternFill("solid", fgColor=STATUS_COLORS[row.status])
                for column_number, cell in enumerate(sheet[sheet.max_row], start=1):
                    cell.fill = fill
                    cell.border = border
                    cell.alignment = Alignment(
                        horizontal="left" if column_number in {1, 2, 8, 10} else "right",
                        vertical="center",
                        wrap_text=column_number in {1, 2, 10},
                    )
                    if 3 <= column_number <= 7:
                        cell.number_format = "0.######"
                    elif column_number == 9:
                        cell.number_format = '0.##"%"'

    if not report.groups:
        sheet.append(["За обраний період немає даних для перевірки."])

    sheet.freeze_panes = "C7"
    for column_number, width in enumerate(
        [28, 34, 16, 16, 22, 16, 16, 12, 16, 28],
        start=1,
    ):
        sheet.column_dimensions[get_column_letter(column_number)].width = width
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:6"

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
